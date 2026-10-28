from decimal import Decimal

import pytest
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import purchase as purchase_service
from app.api.purchase import to_xlsx
from app.database import Base
from app.models.organization_contract import Organization
from app.models.project import Project
from app.models.purchase import PurchaseItem
from app.models.rd import RdSpecItem, RdVolume


@pytest.fixture
def world():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_conn, _rec):
        cur = dbapi_conn.cursor()
        cur.execute("PRAGMA foreign_keys=ON")
        cur.close()

    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="Purchase org")
        db.add(org)
        db.flush()
        project = Project(name="Purchase project", organization_id=org.id)
        other = Project(name="Other", organization_id=org.id)
        db.add_all([project, other])
        db.flush()
        db.commit()
        yield {"db": db, "project": project, "other": other}
    engine.dispose()


def _volume(db, project_id, title="Том 1", ext="v1"):
    vol = RdVolume(
        project_id=project_id,
        title=title,
        file_name=f"{title}.pdf",
        provider="google_drive",
        external_id=ext,
    )
    db.add(vol)
    db.flush()
    return vol


def _spec(db, project_id, volume_id, **kwargs):
    spec = RdSpecItem(
        project_id=project_id,
        volume_id=volume_id,
        status=kwargs.pop("status", "confirmed"),
        **kwargs,
    )
    db.add(spec)
    db.flush()
    return spec


def test_group_key_normalizes():
    a = purchase_service.group_key("Щит ЩР-УПП1 (… – аналогично)", None, "шт")
    b = purchase_service.group_key("щит щр-упп1", None, "шт")
    assert a == b


def test_rebuild_groups_and_sums(world):
    db = world["db"]
    project = world["project"]
    vol1 = _volume(db, project.id, "Том А", "va")
    vol2 = _volume(db, project.id, "Том Б", "vb")
    _spec(db, project.id, vol1.id, name="Кабель ВВГ", unit="м", qty=Decimal("10"))
    _spec(db, project.id, vol2.id, name="кабель ввг", unit="м", qty=Decimal("5"))
    _spec(db, project.id, vol1.id, name="Розетка", unit="шт", qty=Decimal("2"),
          status="proposed")
    _spec(db, project.id, vol1.id, name="Лампа", unit="шт", qty=Decimal("3"),
          status="rejected")
    db.commit()

    res = purchase_service.rebuild(db, project.id)
    assert res["items"] == 1
    assert res["created"] == 1
    assert res["updated"] == 0

    items = purchase_service.list_items(db, project.id)
    assert len(items) == 1
    assert Decimal(items[0]["qty_needed"]) == 15
    assert sorted(items[0]["volumes"]) == sorted(["Том А", "Том Б"])


def test_rebuild_preserves_status_and_removes_need(world):
    db = world["db"]
    project = world["project"]
    vol = _volume(db, project.id, "Том", "v")
    s1 = _spec(db, project.id, vol.id, name="Щит", unit="шт", qty=Decimal("1"))
    s2 = _spec(db, project.id, vol.id, name="Лампа", unit="шт", qty=Decimal("2"))
    db.commit()

    purchase_service.rebuild(db, project.id)
    items = {it["name"]: it for it in purchase_service.list_items(db, project.id)}
    assert set(items) == {"Щит", "Лампа"}

    purchase_service.set_status(db, project.id, items["Щит"]["id"], "ordered",
                                qty_ordered="1")
    purchase_service.rebuild(db, project.id)
    items = {it["name"]: it for it in purchase_service.list_items(db, project.id)}
    assert items["Щит"]["status"] == "ordered"
    assert Decimal(items["Щит"]["qty_ordered"]) == 1

    s1.status = "rejected"
    s2.status = "rejected"
    db.commit()
    res = purchase_service.rebuild(db, project.id)
    items = {it["name"]: it for it in purchase_service.list_items(db, project.id)}
    assert "Щит" in items
    assert "Лампа" not in items
    assert res["removed"] == 1


def test_set_status_rules(world):
    db = world["db"]
    project = world["project"]
    other = world["other"]
    vol = _volume(db, project.id, "Том", "v")
    _spec(db, project.id, vol.id, name="Щит", unit="шт", qty=Decimal("4"))
    db.commit()
    purchase_service.rebuild(db, project.id)
    item = purchase_service.list_items(db, project.id)[0]

    res = purchase_service.set_status(db, project.id, item["id"], "ordered")
    assert res["status"] == "ordered"
    assert Decimal(res["qty_ordered"]) == 4

    res = purchase_service.set_status(db, project.id, item["id"], "need")
    assert Decimal(res["qty_ordered"]) == 0

    from fastapi import HTTPException
    with pytest.raises(HTTPException) as exc:
        purchase_service.set_status(db, other.id, item["id"], "ordered")
    assert exc.value.status_code == 404

    with pytest.raises(HTTPException) as exc2:
        purchase_service.set_status(db, project.id, item["id"], "bogus")
    assert exc2.value.status_code == 422


def test_to_xlsx_roundtrip(world):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    from io import BytesIO

    items = [{
        "name": "Щит", "type_mark": "ЩР", "manufacturer": "M",
        "unit": "шт", "qty_needed": "2", "qty_ordered": "0",
        "qty_left": "2", "status": "need", "note": None,
    }]
    data = to_xlsx(items)
    assert isinstance(data, bytes)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.title == "Закупка"
    headers = [c.value for c in ws[1]]
    assert headers[0] == "№"
    assert headers[1] == "Наименование"
    assert ws.max_row == 2
    assert ws.cell(row=2, column=9).value == "Нужно заказать"


def test_api_endpoints(world, monkeypatch):
    from fastapi import FastAPI, HTTPException
    from fastapi.testclient import TestClient
    from app.api import purchase as purchase_api
    from app.core.auth import require_user
    from app.database import get_db

    db = world["db"]
    project = world["project"]
    vol = _volume(db, project.id, "Том", "v")
    _spec(db, project.id, vol.id, name="Щит", unit="шт", qty=Decimal("1"))
    db.commit()

    fake_user = type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    app = FastAPI()
    app.include_router(purchase_api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_user] = lambda: fake_user
    client = TestClient(app)

    r = client.get(f"/projects/{project.id}/purchase")
    assert r.status_code == 200
    assert "items" in r.json() and "summary" in r.json()

    r = client.post(f"/projects/{project.id}/purchase/rebuild")
    assert r.status_code == 200
    assert r.json()["items"] == 1

    item_id = client.get(f"/projects/{project.id}/purchase").json()["items"][0]["id"]
    r = client.patch(f"/projects/{project.id}/purchase/{item_id}", json={"status": "ordered"})
    assert r.status_code == 200
    assert r.json()["status"] == "ordered"

    r = client.get(f"/projects/{project.id}/purchase/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    def viewer_only(_db, _user, _pid, minimum="viewer"):
        if minimum != "viewer":
            raise HTTPException(status_code=403, detail="Insufficient project access")
        return "viewer"

    monkeypatch.setattr(purchase_api, "require_project_role", viewer_only)
    assert client.get(f"/projects/{project.id}/purchase").status_code == 200
    assert client.post(f"/projects/{project.id}/purchase/rebuild").status_code == 403
    assert client.patch(f"/projects/{project.id}/purchase/{item_id}", json={"status": "paid"}).status_code == 403
