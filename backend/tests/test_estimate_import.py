"""Тесты разбора сметы из Excel (.xlsx)."""

from io import BytesIO

import pytest
from openpyxl import Workbook

from app.estimate_import import parse_estimate_xlsx


def _wb_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_simple_table():
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.append(["№", "Шифр", "Наименование", "Ед. изм.", "Количество", "Цена", "Всего"])
    ws.append([1, "ФЕР01-01-001-01", "Разработка грунта", "м3", 10, 100, 1000])
    ws.append([2, "ФЕР01-01-002-01", "Устройство фундамента", "м3", 5, 200, 1000])
    ws.append([3, "ТЕР01-01-003-01", "Устройство стен", "м2", 20, 50, 1000])

    result = parse_estimate_xlsx(_wb_bytes(wb))

    lines = result["lines"]
    assert len(lines) == 3

    l0 = lines[0]
    assert l0["sort_order"] == 1
    assert l0["position"] == "1"
    assert l0["code"] == "ФЕР01-01-001-01"
    assert l0["name"] == "Разработка грунта"
    assert l0["unit"] == "м3"
    assert l0["qty"] == "10"
    assert l0["unit_price"] == "100"
    assert l0["total"] == "1000"

    l1 = lines[1]
    assert l1["position"] == "2"
    assert l1["name"] == "Устройство фундамента"

    assert result["detected"]["sheet"] == "Смета"
    assert result["detected"]["header_row"] == 0
    assert "name" in result["detected"]["columns"]
    assert "qty" in result["detected"]["columns"]

    # Нет несходимости.
    assert result["warnings"] == []
    assert result["total"] is None


def test_two_row_header_and_sections():
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.append(["", "", "Наименование", "", "Количество", "", ""])
    ws.append(["№", "Шифр", "", "Ед. изм.", "", "Цена", "Всего"])
    ws.append(["Раздел 1. Земляные работы", "", "", "", "", "", ""])
    ws.append([1, "ФЕР", "Разработка грунта", "м3", 10, 100, 1000])
    ws.append([2, "ФЕР", "Обратная засыпка", "м3", 5, 200, 1000])
    ws.append(["Раздел 2. Фундаменты", "", "", "", "", "", ""])
    ws.append([3, "ФЕР", "Устройство фундамента", "м3", 4, 250, 1000])
    ws.append(["Итого", "", "", "", "", "", 3000])
    ws.append(["Итого по смете", "", "", "", "", "", 3000])

    result = parse_estimate_xlsx(_wb_bytes(wb))
    lines = result["lines"]
    assert len(lines) == 3
    assert lines[0]["section"] == "Раздел 1. Земляные работы"
    assert lines[1]["section"] == "Раздел 1. Земляные работы"
    assert lines[2]["section"] == "Раздел 2. Фундаменты"
    assert result["total"] == "3000"

    positions = [l["position"] for l in lines]
    assert positions == ["1", "2", "3"]


def test_unit_in_name_tail():
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Количество", "Цена", "Всего"])
    ws.append([1, "Устройство стяжки, 100 м2", 2, 1000, 2000])

    result = parse_estimate_xlsx(_wb_bytes(wb))
    lines = result["lines"]
    assert len(lines) == 1
    line = lines[0]
    assert line["name"] == "Устройство стяжки"
    assert line["unit"] == "100 м2"


def test_unit_price_recalc_and_warning():
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Ед. изм.", "Количество", "Цена", "Всего"])
    # Цена пустая — должна посчитаться как total/qty = 100.
    ws.append([1, "Позиция A", "шт", 5, None, 500])
    # Сумма не сходится: qty * price = 200, а total = 500 (>1%).
    ws.append([2, "Позиция B", "шт", 2, 100, 500])

    result = parse_estimate_xlsx(_wb_bytes(wb))
    lines = result["lines"]
    assert lines[0]["unit_price"] == "100"
    assert lines[0]["total"] == "500"
    assert any("Позиция 2" in w for w in result["warnings"])


def test_comma_decimals():
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Ед. изм.", "Количество", "Цена", "Всего"])
    ws.append([1, "Позиция", "м3", "10,5", "100,25", "1052,625"])

    result = parse_estimate_xlsx(_wb_bytes(wb))
    lines = result["lines"]
    assert len(lines) == 1
    assert lines[0]["qty"] == "10.5"
    assert lines[0]["unit_price"] == "100.25"


def test_numeric_cells_only():
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Ед. изм.", "Количество", "Цена", "Всего"])
    ws.append(["1", "Позиция", "шт", 3.5, 10.5, 36.75])

    result = parse_estimate_xlsx(_wb_bytes(wb))
    lines = result["lines"]
    assert len(lines) == 1
    line = lines[0]
    assert line["qty"] == "3.5"
    assert line["unit_price"] == "10.5"
    assert line["total"] == "36.75"


def test_no_table_raises():
    wb = Workbook()
    ws = wb.active
    ws.append(["какие-то", "данные"])
    ws.append(["ничего", "полезного"])

    with pytest.raises(ValueError):
        parse_estimate_xlsx(_wb_bytes(wb))


def test_too_many_lines_raises():
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Ед. изм.", "Количество", "Цена", "Всего"])
    for i in range(5001):
        ws.append([i + 1, "Позиция", "шт", 1, 10, 10])

    with pytest.raises(ValueError):
        parse_estimate_xlsx(_wb_bytes(wb))


_SPECIAL_VALUES = ("NaN", "sNaN", "Infinity", "-Infinity")


def _special_wb(qty, price, total) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Количество", "Цена", "Всего"])
    ws.append([1, "Тест", qty, price, total])
    return _wb_bytes(wb)


@pytest.mark.parametrize("special", _SPECIAL_VALUES)
@pytest.mark.parametrize("column", ["qty", "unit_price", "total"])
@pytest.mark.parametrize("case", ["lower", "upper", "mixed"])
def test_special_numeric_values_raise(special, column, case):
    value = {"lower": special.lower(), "upper": special.upper(), "mixed": special}[case]
    cells = {"qty": "1", "unit_price": "50", "total": "100"}
    cells[column] = value
    data = _special_wb(cells["qty"], cells["unit_price"], cells["total"])

    with pytest.raises(ValueError):
        parse_estimate_xlsx(data)


@pytest.mark.parametrize("special", _SPECIAL_VALUES)
def test_special_grand_total_raises(special):
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Количество", "Цена", "Всего"])
    ws.append([1, "Тест", 1, 50, 100])
    ws.append(["Итого по смете", "", "", "", special])

    with pytest.raises(ValueError):
        parse_estimate_xlsx(_wb_bytes(wb))


def test_regular_numbers_control():
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Наименование", "Количество", "Цена", "Всего"])
    ws.append([1, "Тест", 2, 50, 100])

    result = parse_estimate_xlsx(_wb_bytes(wb))
    lines = result["lines"]
    assert len(lines) == 1
    assert lines[0]["qty"] == "2"
    assert lines[0]["unit_price"] == "50"
    assert lines[0]["total"] == "100"
    assert result["warnings"] == []
