"""Тесты протокола проверки цепи «фаза–нуль»."""

import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.phase_zero_protocol import (
    CHARACTERISTIC_MULTIPLIER,
    fault_current,
    required_current,
    to_xlsx,
    validate_rows,
    verdict,
)


def test_validate_rows_accepts_comma_and_skips_blank():
    rows = [
        {
            "circuit_name": "ЩР-1 / Л-1",
            "device_type": "автоматический выключатель",
            "device_mark": "ВА47-29",
            "rated_a": "16",
            "characteristic": "C",
            "phase_voltage_v": "230",
            "loop_ohm": "0,8",
        },
        {"circuit_name": "", "device_type": "", "rated_a": ""},
        {},
        {
            "circuit_name": "ЩР-2 / Л-3",
            "device_type": "предохранитель",
            "rated_a": 25,
            "short_circuit_a": "500",
        },
    ]
    out = validate_rows(rows)
    assert len(out) == 2
    assert out[0]["rated_a"] == Decimal("16")
    assert out[0]["loop_ohm"] == Decimal("0.8")
    assert out[0]["phase_voltage_v"] == Decimal("230")
    assert out[1]["phase_voltage_v"] == Decimal("220")
    assert out[1]["short_circuit_a"] == Decimal("500")
    assert out[1]["device_type"] == "предохранитель"


@pytest.mark.parametrize(
    "row, needle",
    [
        (
            {
                "circuit_name": "Л-1",
                "device_type": "автоматический выключатель",
                "rated_a": 0,
                "characteristic": "C",
            },
            "Строка 1",
        ),
        (
            {
                "circuit_name": "Л-1",
                "device_type": "автоматический выключатель",
                "rated_a": 16,
                "characteristic": "A",
            },
            "Строка 1",
        ),
        (
            {
                "circuit_name": "Л-1",
                "device_type": "неизвестный",
                "rated_a": 16,
            },
            "Строка 1",
        ),
        (
            {
                "circuit_name": "",
                "device_type": "автоматический выключатель",
                "rated_a": 16,
                "characteristic": "C",
            },
            "Строка 1",
        ),
        (
            {
                "circuit_name": "Л-1",
                "device_type": "автоматический выключатель",
                "rated_a": 16,
                "characteristic": "C",
                "loop_ohm": -1,
            },
            "Строка 1",
        ),
    ],
)
def test_validate_rows_errors(row, needle):
    with pytest.raises(ValueError) as exc:
        validate_rows([row])
    assert needle in str(exc.value)


def test_characteristic_multipliers():
    assert CHARACTERISTIC_MULTIPLIER["B"] == Decimal("5")
    assert CHARACTERISTIC_MULTIPLIER["C"] == Decimal("10")
    assert CHARACTERISTIC_MULTIPLIER["D"] == Decimal("20")


def test_fault_current_from_loop():
    row = {"phase_voltage_v": Decimal("220"), "loop_ohm": Decimal("1")}
    assert fault_current(row) == Decimal("220")
    row2 = {"phase_voltage_v": Decimal("220"), "loop_ohm": Decimal("2")}
    assert fault_current(row2) == Decimal("110")
    row3 = {"short_circuit_a": Decimal("300")}
    assert fault_current(row3) == Decimal("300")
    assert fault_current({}) is None


def test_required_current_automat_and_fuse():
    autom = {
        "device_type": "автоматический выключатель",
        "rated_a": Decimal("16"),
        "characteristic": "C",
    }
    # 1.1 * 10 * 16 = 176
    assert required_current(autom) == Decimal("176.0")

    автом_b = {
        "device_type": "автоматический выключатель",
        "rated_a": Decimal("16"),
        "characteristic": "B",
    }
    # 1.1 * 5 * 16 = 88
    assert required_current(автом_b) == Decimal("88.0")

    автом_d = {
        "device_type": "автоматический выключатель",
        "rated_a": Decimal("10"),
        "characteristic": "D",
    }
    # 1.1 * 20 * 10 = 220
    assert required_current(автом_d) == Decimal("220.0")

    fuse = {
        "device_type": "предохранитель",
        "rated_a": Decimal("20"),
    }
    # 1.1 * 3 * 20 = 66
    assert required_current(fuse) == Decimal("66.0")

    uzо = {
        "device_type": "УЗО",
        "rated_a": Decimal("40"),
    }
    assert required_current(uzо) is None


def test_verdict_states():
    # годен
    assert (
        verdict(
            {
                "device_type": "автоматический выключатель",
                "rated_a": Decimal("16"),
                "characteristic": "C",
                "phase_voltage_v": Decimal("220"),
                "loop_ohm": Decimal("1"),
            }
        )
        == "годен"
    )
    # 220 >= 176 -> годен; теперь сделаем негодный: Z=1.5 -> 146.67 < 176
    assert (
        verdict(
            {
                "device_type": "автоматический выключатель",
                "rated_a": Decimal("16"),
                "characteristic": "C",
                "phase_voltage_v": Decimal("220"),
                "loop_ohm": Decimal("1.5"),
            }
        )
        == "не годен"
    )
    # нет данных
    assert (
        verdict(
            {
                "device_type": "автоматический выключатель",
                "rated_a": Decimal("16"),
                "characteristic": "C",
            }
        )
        == "нет данных"
    )
    # УЗО
    assert (
        verdict(
            {
                "device_type": "УЗО",
                "rated_a": Decimal("40"),
                "phase_voltage_v": Decimal("220"),
                "loop_ohm": Decimal("2"),
            }
        )
        == "годен (проверяется кнопкой «Тест»)"
    )
    # УЗО без данных
    assert verdict({"device_type": "УЗО", "rated_a": Decimal("40")}) == "нет данных"


def test_verdict_boundary_equality():
    # 1.1 * 10 * 16 = 176; I = U/Z. Z = 220 / 176 = 1.25
    row = {
        "device_type": "автоматический выключатель",
        "rated_a": Decimal("16"),
        "characteristic": "C",
        "phase_voltage_v": Decimal("220"),
        "loop_ohm": Decimal("1.25"),
    }
    assert fault_current(row) == Decimal("176")
    assert required_current(row) == Decimal("176.0")
    assert verdict(row) == "годен"

    # чуть меньше требуемого
    row2 = {
        "device_type": "автоматический выключатель",
        "rated_a": Decimal("16"),
        "characteristic": "C",
        "phase_voltage_v": Decimal("220"),
        "loop_ohm": Decimal("1.26"),
    }
    assert verdict(row2) == "не годен"


def test_to_xlsx_layout_and_red_fill():
    header = {
        "object_name": "Объект",
        "protocol_no": "7",
        "date": "2025-02-02",
        "tested_by": "Иванов",
        "checked_by": "Петров",
    }
    rows = [
        {
            "circuit_name": "ЩР-1 / Л-1",
            "device_type": "автоматический выключатель",
            "device_mark": "ВА47-29",
            "rated_a": 16,
            "characteristic": "C",
            "phase_voltage_v": 220,
            "loop_ohm": 1,
        },
        {
            "circuit_name": "ЩР-1 / Л-2",
            "device_type": "автоматический выключатель",
            "device_mark": "ВА47-29",
            "rated_a": 16,
            "characteristic": "C",
            "phase_voltage_v": 220,
            "loop_ohm": 2,
        },
    ]
    data = to_xlsx(header, rows)
    assert isinstance(data, bytes) and data[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Протокол"]

    found_title = False
    found_conclusion = False
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                if "ПРОТОКОЛ № 7" in cell.value:
                    found_title = True
                if cell.value.startswith("Заключение:"):
                    found_conclusion = True
        if any(isinstance(c.value, str) and c.value == "№ п/п" for c in row):
            header_row = row[0].row
    assert found_title
    assert found_conclusion
    assert header_row is not None

    headers = [c.value for c in ws[header_row]]
    assert "Цепь" in headers
    assert "Iкз А" in headers
    assert "Требуется Iкз ≥ А" in headers
    assert "Заключение" in headers

    concl_idx = headers.index("Заключение") + 1
    red_found = False
    for r in range(header_row + 1, header_row + 1 + len(rows)):
        cell = ws.cell(row=r, column=concl_idx)
        if cell.value == "не годен":
            assert cell.fill is not None
            assert cell.fill.fgColor.rgb in ("00FFC7CE", "FFC7CE")
            red_found = True
    assert red_found


class _FakeUser:
    id = 1
    username = "tester"
    is_admin = False


def _client_for(monkeypatch):
    from app.api import phase_zero_protocol as api_mod

    app = FastAPI()
    app.include_router(api_mod.router)

    def _fake_require_user():
        return _FakeUser()

    def _fake_require_project_role(db, user, project_id, role="viewer"):
        return None

    def _fake_get_db():
        yield None

    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "require_project_role", _fake_require_project_role)
    monkeypatch.setattr(api_mod, "get_db", lambda: (yield None))
    return TestClient(app)


def test_api_returns_xlsx(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1", "object_name": "О"},
        "rows": [
            {
                "circuit_name": "Л-1",
                "device_type": "автоматический выключатель",
                "rated_a": 16,
                "characteristic": "C",
                "phase_voltage_v": 220,
                "loop_ohm": 1,
            },
        ],
    }
    resp = client.post(
        "/projects/1/executive-docs/phase-zero-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert "attachment" in resp.headers.get("content-disposition", "")


def test_api_returns_422_on_error(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [
            {
                "circuit_name": "Л-1",
                "device_type": "автоматический выключатель",
                "rated_a": 0,
                "characteristic": "C",
            },
        ],
    }
    resp = client.post(
        "/projects/1/executive-docs/phase-zero-protocol.xlsx", json=body,
    )
    assert resp.status_code == 422
    assert "Строка 1" in resp.json()["detail"]


def test_api_viewer_sufficient(monkeypatch):
    from app.api import phase_zero_protocol as api_mod

    calls = {}

    def _fake_get_db():
        yield None

    def _fake_require_user():
        return _FakeUser()

    def _fake_role(db, user, project_id, role="viewer"):
        calls["role"] = role
        return None

    app = FastAPI()
    app.include_router(api_mod.router)
    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "get_db", _fake_get_db)
    monkeypatch.setattr(api_mod, "require_project_role", _fake_role)

    client = TestClient(app)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [
            {
                "circuit_name": "Л-1",
                "device_type": "предохранитель",
                "rated_a": 20,
                "short_circuit_a": 100,
            },
        ],
    }
    resp = client.post(
        "/projects/1/executive-docs/phase-zero-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert calls.get("role") == "viewer"
