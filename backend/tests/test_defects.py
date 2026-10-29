from datetime import date, timedelta
from io import BytesIO

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import defects as svc
from app.database import Base
from app.models.organization_contract import Organization
from app.models.project import Project

JPEG = b"\xff\xd8\xff\xe0 fake jpeg"


@pytest.fixture
def world():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="Def org")
        db.add(org)
        db.flush()
        p = Project(name="P", organization_id=org.id)
        o = Project(name="O", organization_id=org.id)
        db.add_all([p, o])
        db.flush()
        db.commit()
        yield {"db": db, "p": p, "o": o}
    engine.dispose()


def test_numbers_are_per_project_and_validation(world):
    db, pid = world["db"], world["p"].id
    a = svc.create_defect(db, pid, None, {"title": "Трещина в стяжке", "location": "Этаж 2", "priority": "high"})
    b = svc.create_defect(db, pid, None, {"title": "Нет заземления"})
    c = svc.create_defect(db, world["o"].id, None, {"title": "Другой проект"})
    assert (a["number"], b["number"], c["number"]) == (1, 2, 1)
    assert a["status"] == "open" and a["priority"] == "high"
    with pytest.raises(ValueError):
        svc.create_defect(db, pid, None, {"title": "  "})
    with pytest.raises(ValueError):
        svc.create_defect(db, pid, None, {"title": "x", "priority": "urgent"})
    with pytest.raises(ValueError):
        svc.create_defect(db, pid, None, {"title": "x", "due_date": "завтра"})


def test_status_flow_and_fix_needs_evidence(world):
    db, pid = world["db"], world["p"].id
    d = svc.create_defect(db, pid, None, {"title": "Не закреплён лоток"})
    with pytest.raises(ValueError):
        svc.change_status(db, pid, None, d["id"], "accepted", None)  # из «Открыто» сразу в «Принято» нельзя
    svc.change_status(db, pid, None, d["id"], "in_progress", "Взяли в работу")
    with pytest.raises(ValueError):
        svc.change_status(db, pid, None, d["id"], "fixed", " ")  # нужен результат: фото или комментарий
    svc.add_photo(db, pid, d["id"], "after", "после.jpg", "image/jpeg", JPEG)
    fixed = svc.change_status(db, pid, None, d["id"], "fixed", None)
    assert fixed["status"] == "fixed"
    accepted = svc.change_status(db, pid, None, d["id"], "accepted", "Проверил")
    assert accepted["status"] == "accepted" and accepted["closed_at"]
    assert [c["status_change"] for c in accepted["comments"]] == ["in_progress", "fixed", "accepted"]
    reopened = svc.change_status(db, pid, None, d["id"], "open", "Опять отвалился")
    assert reopened["status"] == "open" and reopened["closed_at"] is None
    with pytest.raises(ValueError):
        svc.change_status(db, pid, None, d["id"], "open", None)


def test_photos_rules_and_isolation(world):
    db, pid = world["db"], world["p"].id
    d = svc.create_defect(db, pid, None, {"title": "Т"})
    res = svc.add_photo(db, pid, d["id"], "before", "до.jpg", "image/jpeg; charset=x", JPEG)
    assert res["photos"][0]["kind"] == "before" and res["photos"][0]["size_bytes"] == len(JPEG)
    for args in (("before", "a.gif", "image/gif", JPEG), ("before", "a.jpg", "image/jpeg", b""), ("maybe", "a.jpg", "image/jpeg", JPEG)):
        with pytest.raises(ValueError):
            svc.add_photo(db, pid, d["id"], *args)
    with pytest.raises(ValueError):
        svc.add_photo(db, pid, d["id"], "before", "big.jpg", "image/jpeg", b"x" * (svc.MAX_PHOTO_BYTES + 1))
    for i in range(svc.MAX_PHOTOS - 1):
        svc.add_photo(db, pid, d["id"], "before", f"{i}.jpg", "image/jpeg", JPEG)
    with pytest.raises(ValueError):
        svc.add_photo(db, pid, d["id"], "before", "last.jpg", "image/jpeg", JPEG)
    with pytest.raises(LookupError):
        svc.add_photo(db, world["o"].id, d["id"], "before", "x.jpg", "image/jpeg", JPEG)
    photo_id = svc.get_defect(db, pid, d["id"])["photos"][0]["id"]
    assert svc.get_photo(db, pid, d["id"], photo_id).data == JPEG
    with pytest.raises(LookupError):
        svc.get_photo(db, world["o"].id, d["id"], photo_id)
    assert len(svc.delete_photo(db, pid, d["id"], photo_id)["photos"]) == svc.MAX_PHOTOS - 1


def test_list_filters_counts_overdue_and_export(world):
    from openpyxl import load_workbook

    db, pid = world["db"], world["p"].id
    old = (date.today() - timedelta(days=3)).isoformat()
    a = svc.create_defect(db, pid, None, {"title": "Просроченное", "due_date": old, "location": "Этаж 3"})
    svc.create_defect(db, pid, None, {"title": "Свежее"})
    svc.change_status(db, pid, None, a["id"], "in_progress", None)
    res = svc.list_defects(db, pid)
    assert res["counts"]["open"] == 1 and res["counts"]["in_progress"] == 1 and res["overdue"] == 1
    assert [d["title"] for d in svc.list_defects(db, pid, status="in_progress")["defects"]] == ["Просроченное"]
    assert [d["title"] for d in svc.list_defects(db, pid, q="ЭТАЖ 3")["defects"]] == ["Просроченное"]
    ws = load_workbook(BytesIO(svc.to_xlsx(db, pid)))["Замечания"]
    assert ws["A1"].value == "№" and ws.max_row == 3 and ws["J2"].value == "да"
    svc.update_defect(db, pid, a["id"], {"responsible": "Иванов", "due_date": None, "title": "Новое имя"})
    assert svc.get_defect(db, pid, a["id"])["responsible"] == "Иванов"
    with pytest.raises(ValueError):
        svc.update_defect(db, pid, a["id"], {"title": " "})
    svc.delete_defect(db, pid, a["id"])
    assert svc.list_defects(db, pid)["counts"]["in_progress"] == 0
    with pytest.raises(LookupError):
        svc.get_defect(db, pid, a["id"])


def test_api(world):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api import defects as api
    from app.core.auth import require_user
    from app.database import get_db

    db, pid = world["db"], world["p"].id
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_user] = lambda: type("U", (), {"id": None, "is_active": True, "is_admin": True})()
    c = TestClient(app)
    r = c.post(f"/projects/{pid}/defects", json={"title": "Течь", "location": "Подвал"})
    assert r.status_code == 200
    did = r.json()["id"]
    assert c.post(f"/projects/{pid}/defects", json={"title": " "}).status_code == 422
    r = c.post(f"/projects/{pid}/defects/{did}/photos", params={"name": "a.jpg"}, content=JPEG, headers={"content-type": "image/jpeg"})
    assert r.status_code == 200
    photo = r.json()["photos"][0]["id"]
    assert c.get(f"/projects/{pid}/defects/{did}/photos/{photo}").content == JPEG
    assert c.post(f"/projects/{pid}/defects/{did}/photos", content=b"x", headers={"content-type": "text/plain"}).status_code == 422
    assert c.post(f"/projects/{pid}/defects/{did}/comments", json={"text": "Сфотографировал"}).json()["comments"][0]["text"] == "Сфотографировал"
    assert c.post(f"/projects/{pid}/defects/{did}/status", json={"status": "accepted"}).status_code == 422
    assert c.post(f"/projects/{pid}/defects/{did}/status", json={"status": "in_progress"}).json()["status"] == "in_progress"
    assert c.get(f"/projects/{pid}/defects").json()["counts"]["in_progress"] == 1
    assert c.get(f"/projects/{pid}/defects/export.xlsx").content[:2] == b"PK"
    assert c.get(f"/projects/{world['o'].id}/defects/{did}").status_code == 404
    assert c.delete(f"/projects/{pid}/defects/{did}").json() == {"deleted": did}
