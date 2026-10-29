"""Тесты протокола измерения сопротивления заземляющего устройства."""

import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.earthing_protocol import (
    norm_ohm,
    reduced_ohm,
    row_norm,
    to_xlsx,
    validate_rows,
    verdict,
)


def test_validate_rows_accepts_comma_and_skips_blank():
    rows = [
        {"point_name": "ЗУ-1", "location": "ТП-1", "soil": "суглинок",
         "electrode": "уголок 50х50х5", "network_voltage_v": 380,
         "measured_ohm": "3,5", "season_factor": "1,2"},
        {"point_name": "", "measured_ohm": ""},
        {},
        {"point_name": "ЗУ-2", "network_voltage_v": 220, "measured_ohm": "6"},
    ]
    out = validate_rows(rows)
    assert len(out) == 2
    assert out[0]["measured_ohm"] == Decimal("3.50")
    assert out[0]["season_factor"] == Decimal("1.2")
    assert out[1]["network_voltage_v"] == 220
    assert out[1]["season_factor"] == Decimal("1.0")


@pytest.mark.parametrize(
    "row, needle",
    [
        ({"point_name": "ЗУ-1", "measured_ohm": -1}, "Строка 1"),
        ({"point_name": "ЗУ-1", "network_voltage_v": 400}, "Строка 1"),
        ({"point_name": "ЗУ-1", "season_factor": 0.5}, "Строка 1"),
        ({"measured_ohm": 5}, "Строка 1"),
    ],
)
def test_validate_rows_errors(row, needle):
    with pytest.raises(ValueError) as exc:
        validate_rows([row])
    assert needle in str(exc.value)


def test_norm_ohm_by_voltage():
    assert norm_ohm(380) == Decimal("4")
    assert norm_ohm(220) == Decimal("8")
    assert norm_ohm(660) == Decimal("2")
    with pytest.raises(ValueError):
        norm_ohm(127)


def test_row_norm_override():
    row = {"network_voltage_v": 380, "norm_override": Decimal("10")}
    assert row_norm(row) == Decimal("10")
    assert row_norm({"network_voltage_v": 380}) == Decimal("4")


def test_reduced_ohm():
    row = {"measured_ohm": Decimal("3.5"), "season_factor": Decimal("1.2")}
    assert reduced_ohm(row) == Decimal("4.20")
    assert reduced_ohm({"measured_ohm": None}) is None


def test_verdict_states():
    assert verdict({"network_voltage_v": 380,
                    "measured_ohm": Decimal("3"),
                    "season_factor": Decimal("1")}) == "годен"
    assert verdict({"network_voltage_v": 380,
                    "measured_ohm": Decimal("4"),
                    "season_factor": Decimal("1")}) == "годен"
    assert verdict({"network_voltage_v": 380,
                    "measured_ohm": Decimal("5"),
                    "season_factor": Decimal("1")}) == "не годен"
    assert verdict({"network_voltage_v": 380,
                    "measured_ohm": Decimal("3"),
                    "season_factor": Decimal("1.5")}) == "не годен"
    assert verdict({"network_voltage_v": 380,
                    "measured_ohm": None,
                    "season_factor": Decimal("1")}) == "нет данных"


def test_to_xlsx_layout_and_red_fill():
    header = {
        "object_name": "Объект",
        "protocol_no": "7",
        "date": "2025-01-01",
        "tested_by": "Иванов",
        "checked_by": "Петров",
    }
    rows = [
        {"point_name": "ЗУ-1", "location": "ТП-1", "soil": "суглинок",
         "electrode": "уголок", "network_voltage_v": 380,
         "measured_ohm": 3, "season_factor": 1},
        {"point_name": "ЗУ-2", "location": "ТП-2", "soil": "песок",
         "electrode": "труба", "network_voltage_v": 380,
         "measured_ohm": 5, "season_factor": 1},
    ]
    data = to_xlsx(header, rows)
    assert isinstance(data, bytes) and data[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Протокол"]

    found_title = False
    found_conclusion = False
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                if "ПРОТОКОЛ № 7" in cell.value:
                    found_title = True
                if cell.value.startswith("Заключение:"):
                    found_conclusion = True
        if any(isinstance(c.value, str) and c.value == "№ п/п" for c in row):
            header_row = row[0].row
    assert found_title
    assert found_conclusion
    assert header_row is not None

    headers = [c.value for c in ws[header_row]]
    assert "Точка/ЗУ" in headers
    assert "Измерено Ом" in headers
    assert "Приведённое Ом" in headers
    assert "Норма Ом" in headers

    measured_idx = headers.index("Измерено Ом") + 1
    red_found = False
    for r in range(header_row + 1, header_row + 1 + len(rows)):
        cell = ws.cell(row=r, column=measured_idx)
        if cell.value is not None and float(cell.value) == pytest.approx(5.0):
            assert cell.fill is not None
            assert cell.fill.fgColor.rgb in ("00FFC7CE", "FFC7CE")
            red_found = True
    assert red_found


class _FakeUser:
    id = 1
    username = "tester"
    is_admin = False


def _client_for(monkeypatch):
    from app.api import earthing_protocol as api_mod

    app = FastAPI()
    app.include_router(api_mod.router)

    def _fake_require_user():
        return _FakeUser()

    def _fake_require_project_role(db, user, project_id, role="viewer"):
        return None

    def _fake_get_db():
        yield None

    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "require_project_role", _fake_require_project_role)
    monkeypatch.setattr(api_mod, "get_db", lambda: (yield None))
    return TestClient(app)


def test_api_returns_xlsx(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1", "object_name": "О"},
        "rows": [
            {"point_name": "ЗУ-1", "measured_ohm": 3, "network_voltage_v": 380},
        ],
    }
    resp = client.post(
        "/projects/1/executive-docs/earthing-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert "attachment" in resp.headers.get("content-disposition", "")


def test_api_returns_422_on_error(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [{"point_name": "ЗУ-1", "measured_ohm": -1}],
    }
    resp = client.post(
        "/projects/1/executive-docs/earthing-protocol.xlsx", json=body,
    )
    assert resp.status_code == 422
    assert "Строка 1" in resp.json()["detail"]


def test_api_viewer_sufficient(monkeypatch):
    from app.api import earthing_protocol as api_mod

    calls = {}

    def _fake_get_db():
        yield None

    def _fake_require_user():
        return _FakeUser()

    def _fake_role(db, user, project_id, role="viewer"):
        calls["role"] = role
        return None

    app = FastAPI()
    app.include_router(api_mod.router)
    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "get_db", _fake_get_db)
    monkeypatch.setattr(api_mod, "require_project_role", _fake_role)

    client = TestClient(app)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [{"point_name": "ЗУ-1", "measured_ohm": 3}],
    }
    resp = client.post(
        "/projects/1/executive-docs/earthing-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert calls.get("role") == "viewer"
