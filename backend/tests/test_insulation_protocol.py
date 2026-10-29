"""Тесты протокола измерения сопротивления изоляции."""

import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.insulation_protocol import (
    norm_min_mohm,
    to_xlsx,
    validate_rows,
    verdict,
)


def test_validate_rows_accepts_comma_and_skips_blank():
    rows = [
        {"line_name": "Л-1", "cable_mark": "ВВГ 3х2,5", "length_m": "12,5",
         "voltage_v": 1000, "pairs": {"A-B": "10,5", "B-C": "12"}},
        {"line_name": "", "pairs": {}},
        {},
        {"line_name": "Л-2", "voltage_v": 2500, "pairs": {"жила-оболочка": "5"}},
    ]
    out = validate_rows(rows)
    assert len(out) == 2
    assert out[0]["length_m"] == Decimal("12.5")
    assert out[0]["pairs"]["A-B"] == Decimal("10.5")
    assert out[1]["voltage_v"] == 2500
    assert out[1]["length_m"] == Decimal("0")


@pytest.mark.parametrize(
    "row, needle",
    [
        ({"line_name": "Л-1", "pairs": {"A-B": -1}}, "Строка 1"),
        ({"line_name": "Л-1", "pairs": {"X-Y": 5}}, "Строка 1"),
        ({"pairs": {"A-B": 5}}, "Строка 1"),
    ],
)
def test_validate_rows_errors(row, needle):
    with pytest.raises(ValueError) as exc:
        validate_rows([row])
    assert needle in str(exc.value)


def test_norm_min_mohm():
    assert norm_min_mohm(500) == Decimal("0.5")
    assert norm_min_mohm(1000) == Decimal("0.5")
    assert norm_min_mohm(2500) == Decimal("1.0")


def test_verdict_states():
    assert verdict({"voltage_v": 1000, "pairs": {"A-B": Decimal("1")}}) == "годен"
    assert verdict({"voltage_v": 1000, "pairs": {"A-B": Decimal("0.1")}}) == "не годен"
    assert verdict({"voltage_v": 2500, "pairs": {"A-B": Decimal("0.5")}}) == "не годен"
    assert verdict({"voltage_v": 1000, "pairs": {}}) == "нет данных"


def test_to_xlsx_layout_and_red_fill():
    header = {
        "object_name": "Объект",
        "protocol_no": "42",
        "date": "2025-01-01",
        "tested_by": "Иванов",
        "checked_by": "Петров",
    }
    rows = [
        {"line_name": "Л-1", "cable_mark": "ВВГ", "length_m": 10,
         "voltage_v": 1000, "pairs": {"A-B": 5, "B-C": 0.2}},
        {"line_name": "Л-2", "cable_mark": "ВВГ", "length_m": 20,
         "voltage_v": 1000, "pairs": {"A-B": 5}},
    ]
    data = to_xlsx(header, rows)
    assert isinstance(data, bytes) and data[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Протокол"]

    found_title = False
    found_conclusion = False
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                if "ПРОТОКОЛ № 42" in cell.value:
                    found_title = True
                if cell.value.startswith("Заключение:"):
                    found_conclusion = True
        if any(isinstance(c.value, str) and c.value == "№ п/п" for c in row):
            header_row = row[0].row
    assert found_title
    assert found_conclusion
    assert header_row is not None

    headers = [c.value for c in ws[header_row]]
    assert "A-B" in headers
    assert "B-C" in headers
    assert "A-N" not in headers
    assert "жила-оболочка" not in headers

    # Найти негодную ячейку (B-C, значение 0.2) и проверить красную заливку
    bc_idx = headers.index("B-C") + 1
    header_cells = list(ws[header_row])
    red_found = False
    for r in range(header_row + 1, header_row + 1 + len(rows)):
        cell = ws.cell(row=r, column=bc_idx)
        if cell.value is not None and float(cell.value) == pytest.approx(0.2):
            assert cell.fill is not None
            assert cell.fill.fgColor.rgb in ("00FFC7CE", "FFC7CE")
            red_found = True
    assert red_found
    assert header_cells  # sanity


class _FakeUser:
    id = 1
    username = "tester"
    is_admin = False


def _client_for(monkeypatch):
    from app.api import insulation_protocol as api_mod

    app = FastAPI()
    app.include_router(api_mod.router)

    def _fake_require_user():
        return _FakeUser()

    def _fake_require_project_role(db, user, project_id, role="viewer"):
        return None

    def _fake_get_db():
        yield None

    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "require_project_role", _fake_require_project_role)
    monkeypatch.setattr(api_mod, "get_db", lambda: (yield None))
    return TestClient(app)


def test_api_returns_xlsx(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1", "object_name": "О"},
        "rows": [
            {"line_name": "Л-1", "pairs": {"A-B": 5}, "voltage_v": 1000},
        ],
    }
    resp = client.post(
        "/projects/1/executive-docs/insulation-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert "attachment" in resp.headers.get("content-disposition", "")


def test_api_returns_422_on_error(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [{"line_name": "Л-1", "pairs": {"A-B": -1}}],
    }
    resp = client.post(
        "/projects/1/executive-docs/insulation-protocol.xlsx", json=body,
    )
    assert resp.status_code == 422
    assert "Строка 1" in resp.json()["detail"]


def test_api_viewer_sufficient(monkeypatch):
    from app.api import insulation_protocol as api_mod

    calls = {}

    def _fake_get_db():
        yield None

    def _fake_require_user():
        return _FakeUser()

    def _fake_role(db, user, project_id, role="viewer"):
        calls["role"] = role
        return None

    app = FastAPI()
    app.include_router(api_mod.router)
    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "get_db", _fake_get_db)
    monkeypatch.setattr(api_mod, "require_project_role", _fake_role)

    client = TestClient(app)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [{"line_name": "Л-1", "pairs": {"A-B": 5}}],
    }
    resp = client.post(
        "/projects/1/executive-docs/insulation-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert calls.get("role") == "viewer"
