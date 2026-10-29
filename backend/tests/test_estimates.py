import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import estimates as svc
from app import ks_forms
from app import work_log as wl
from app.database import Base
from app.models.organization_contract import Organization
from app.models.project import Project

PARSED = {
    "lines": [
        {"sort_order": 1, "section": "Раздел 1. Электромонтаж", "position": "1", "code": "ФЕР08-02-001", "name": "Прокладка кабеля", "unit": "м",
         "qty": "100", "unit_price": "50.00", "total": "5000.00", "note": ""},
        {"sort_order": 2, "section": "Раздел 1. Электромонтаж", "position": "2", "code": "", "name": "Монтаж щита", "unit": "шт",
         "qty": "2", "unit_price": None, "total": "4000.00", "note": ""},
        {"sort_order": 3, "section": "Раздел 2. Пусконаладка", "position": "3", "code": "", "name": "Испытания", "unit": "компл",
         "qty": "1", "unit_price": "1000", "total": "1000", "note": ""},
    ],
    "total": "10000.00",
}


@pytest.fixture
def world():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="Est org")
        db.add(org)
        db.flush()
        p = Project(name="P", organization_id=org.id)
        o = Project(name="O", organization_id=org.id)
        db.add_all([p, o])
        db.flush()
        est = svc.create_estimate(db, p.id, "Смета 1", PARSED, source_name="s.xlsx")
        db.commit()
        lines = svc.get_estimate(db, p.id, est["id"])["lines"]
        yield {"db": db, "p": p, "o": o, "est": est, "lines": lines}
    engine.dispose()


def _log(db, pid, line_id, day, qty):
    return wl.add_entry(db, pid, None, {"work_date": day, "work_name": "Работа", "qty": qty, "unit": "м", "estimate_line_id": line_id})


def test_create_and_read_estimate(world):
    db, pid = world["db"], world["p"].id
    assert world["est"]["lines_count"] == 3 and world["est"]["total"] == "10000.00" and world["est"]["vat_rate"] == "22.00"
    assert [l["name"] for l in world["lines"]] == ["Прокладка кабеля", "Монтаж щита", "Испытания"]
    assert len(svc.list_estimates(db, pid)) == 1 and svc.list_estimates(db, world["o"].id) == []
    with pytest.raises(ValueError):
        svc.create_estimate(db, pid, " ", PARSED)
    with pytest.raises(ValueError):
        svc.create_estimate(db, pid, "x", {"lines": []})
    with pytest.raises(LookupError):
        svc.get_estimate(db, world["o"].id, world["est"]["id"])


def test_journal_entry_must_use_own_estimate_line(world):
    db = world["db"]
    other_est = svc.create_estimate(db, world["o"].id, "Чужая", PARSED)
    other_line = svc.get_estimate(db, world["o"].id, other_est["id"])["lines"][0]["id"]
    with pytest.raises(LookupError):
        _log(db, world["p"].id, other_line, "2026-10-10", "5")
    sug = wl.suggestions(db, world["p"].id)
    assert [l["name"] for l in sug["estimate_lines"]][:2] == ["Прокладка кабеля", "Монтаж щита"]


def test_progress_and_ks_act_from_fact_with_price_from_total(world):
    db, pid = world["db"], world["p"].id
    cable, shield = world["lines"][0]["id"], world["lines"][1]["id"]
    _log(db, pid, cable, "2026-10-05", "40")
    _log(db, pid, cable, "2026-10-20", "30")
    _log(db, pid, shield, "2026-10-21", "1")
    db.commit()
    est = svc.get_estimate(db, pid, world["est"]["id"])
    row = next(l for l in est["lines"] if l["id"] == cable)
    assert row["fact_qty"] == "70.000" and row["done_percent"] == 70.0
    assert est["done_amount"] == "5500.00"  # 70*50 + 1*2000
    prev = svc.preview_act(db, pid, world["est"]["id"], date(2026, 10, 1), date(2026, 10, 31))
    assert prev["total_ex_vat"] == "5500.00" and len(prev["lines"]) == 2
    shield_line = next(l for l in prev["lines"] if l["line_id"] == shield)
    assert shield_line["unit_price"] == "2000.00"  # цена из стоимости / объём
    act = svc.create_act(db, pid, world["est"]["id"], date(2026, 10, 1), date(2026, 10, 31))
    assert act["number"] == "1" and act["total_ex_vat"] == "5500.00"


def test_second_act_does_not_bill_the_same_volume_and_caps_at_estimate(world):
    db, pid = world["db"], world["p"].id
    cable = world["lines"][0]["id"]
    _log(db, pid, cable, "2026-10-05", "80")
    db.commit()
    svc.create_act(db, pid, world["est"]["id"], date(2026, 10, 1), date(2026, 10, 15))
    _log(db, pid, cable, "2026-10-20", "50")  # в смете всего 100, уже выставлено 80
    db.commit()
    prev = svc.preview_act(db, pid, world["est"]["id"], date(2026, 10, 16), date(2026, 10, 28))
    assert prev["lines"][0]["qty"] == "20.000" and prev["total_ex_vat"] == "1000.00"
    assert any("остаток сметного объёма" in w for w in prev["warnings"])
    act2 = svc.create_act(db, pid, world["est"]["id"], date(2026, 10, 16), date(2026, 10, 28))
    assert act2["number"] == "2"
    data = svc.act_with_totals(db, pid, act2["id"])
    assert data["period_total"] == "1000.00" and data["from_start"] == "5000.00" and data["from_year_start"] == "5000.00"
    assert data["vat"] == "220.00" and data["total_with_vat"] == "1220.00"


def test_empty_period_and_status_rules(world):
    db, pid = world["db"], world["p"].id
    with pytest.raises(ValueError):
        svc.create_act(db, pid, world["est"]["id"], date(2026, 10, 1), date(2026, 10, 31))
    with pytest.raises(ValueError):
        svc.preview_act(db, pid, world["est"]["id"], date(2026, 10, 31), date(2026, 10, 1))
    _log(db, pid, world["lines"][0]["id"], "2026-10-05", "10")
    db.commit()
    act = svc.create_act(db, pid, world["est"]["id"], date(2026, 10, 1), date(2026, 10, 31))
    assert svc.set_act_status(db, pid, act["id"], "signed")["status"] == "signed"
    with pytest.raises(ValueError):
        svc.delete_act(db, pid, act["id"])
    with pytest.raises(ValueError):
        svc.set_act_status(db, pid, act["id"], "weird")


def test_line_edit_validation(world):
    db, pid = world["db"], world["p"].id
    lid = world["lines"][0]["id"]
    svc.update_line(db, pid, lid, {"unit_price": "55,5", "qty": "120"})
    line = next(l for l in svc.get_estimate(db, pid, world["est"]["id"])["lines"] if l["id"] == lid)
    assert line["unit_price"] == "55.50" and line["qty"] == "120.000"
    with pytest.raises(ValueError):
        svc.update_line(db, pid, lid, {"qty": "abc"})
    with pytest.raises(ValueError):
        svc.update_line(db, pid, lid, {"unit_price": "-1"})
    with pytest.raises(LookupError):
        svc.update_line(db, world["o"].id, lid, {"name": "x"})


def test_forms_render(world):
    from openpyxl import load_workbook

    db, pid = world["db"], world["p"].id
    _log(db, pid, world["lines"][0]["id"], "2026-10-05", "40")
    db.commit()
    act = svc.create_act(db, pid, world["est"]["id"], date(2026, 10, 1), date(2026, 10, 31))
    data = svc.act_with_totals(db, pid, act["id"])
    header = {"customer": "ФКУ Налог-Сервис", "contractor": "ООО Монтаж", "object_name": "ЦОД Дубна", "contract_no": "12", "contract_date": "2026-01-10"}
    for blob in (ks_forms.render_ks2_docx(header, data), ks_forms.render_ks3_docx(header, data)):
        with zipfile.ZipFile(io.BytesIO(blob)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        assert 'w:orient="landscape"' in xml
    with zipfile.ZipFile(io.BytesIO(ks_forms.render_ks2_docx(header, data))) as zf:
        x2 = zf.read("word/document.xml").decode("utf-8")
    for needle in ("КС-2", "0322005", "ЦОД Дубна", "Прокладка кабеля", "2 000,00", "ФКУ Налог-Сервис"):
        assert needle in x2, needle
    with zipfile.ZipFile(io.BytesIO(ks_forms.render_ks3_docx(header, data))) as zf:
        x3 = zf.read("word/document.xml").decode("utf-8")
    for needle in ("КС-3", "0322001", "Раздел 1. Электромонтаж", "НДС 22%", "2 440,00"):
        assert needle in x3, needle
    wb = load_workbook(io.BytesIO(ks_forms.render_xlsx(header, data)))
    assert wb.sheetnames == ["КС-2", "КС-3"]


def _xlsx_bytes():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Шифр", "Наименование", "Ед. изм.", "Количество", "Цена", "Всего"])
    ws.append([1, "ФЕР08", "Прокладка кабеля", "м", 100, 50, 5000])
    ws.append([2, "", "Монтаж щита", "шт", 2, 2000, 4000])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_api_import_ks_flow(world):
    from fastapi import FastAPI, HTTPException
    from fastapi.testclient import TestClient
    from app.api import estimates as api
    from app.core.auth import require_user
    from app.database import get_db

    db, pid = world["db"], world["p"].id
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_user] = lambda: type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    c = TestClient(app)

    r = c.post(f"/projects/{pid}/estimates/import", params={"name": "Смета", "preview": "true"}, content=_xlsx_bytes())
    assert r.status_code == 200 and r.json()["lines_total"] == 2 and r.json()["detected"]["columns"]["name"] == 2
    assert len(c.get(f"/projects/{pid}/estimates").json()["estimates"]) == 1  # preview ничего не сохранил (кроме общей сметы из фикстуры)
    r = c.post(f"/projects/{pid}/estimates/import", params={"name": "Смета из Excel", "vat_rate": "20"}, content=_xlsx_bytes())
    assert r.status_code == 200 and r.json()["lines_count"] == 2 and r.json()["vat_rate"] == "20.00"
    eid = r.json()["id"]
    assert c.post(f"/projects/{pid}/estimates/import", params={"name": "x"}, content=b"").status_code == 422
    assert c.post(f"/projects/{pid}/estimates/import", params={"name": "x"}, content=b"not an excel file").status_code == 422

    detail = c.get(f"/projects/{pid}/estimates/{eid}").json()
    line_id = detail["lines"][0]["id"]
    assert c.post(f"/projects/{pid}/work-log", json={}).status_code in (404, 405, 422)  # маршрут журнала не подключён в этом мини-приложении
    wl.add_entry(db, pid, None, {"work_date": "2026-10-05", "work_name": "Кабель", "qty": "60", "unit": "м", "estimate_line_id": line_id})
    db.commit()
    prev = c.get(f"/projects/{pid}/estimates/{eid}/ks-preview", params={"period_from": "2026-10-01", "period_to": "2026-10-28"}).json()
    assert prev["total_ex_vat"] == "3000.00"
    r = c.post(f"/projects/{pid}/estimates/{eid}/ks-acts", json={"period_from": "2026-10-01", "period_to": "2026-10-28"})
    assert r.status_code == 200
    aid = r.json()["id"]
    for path, magic in (("ks2.docx", b"PK"), ("ks3.docx", b"PK"), ("ks.xlsx", b"PK")):
        rr = c.post(f"/projects/{pid}/ks-acts/{aid}/{path}", json={"header": {"object_name": "О"}})
        assert rr.status_code == 200 and rr.content[:2] == magic, path
    assert c.patch(f"/projects/{pid}/ks-acts/{aid}", json={"status": "signed"}).json()["status"] == "signed"
    assert c.delete(f"/projects/{pid}/ks-acts/{aid}").status_code == 422
    assert c.get(f"/projects/{world['o'].id}/estimates/{eid}").status_code == 404
    assert c.delete(f"/projects/{pid}/estimates/{eid}").json() == {"deleted": eid}
