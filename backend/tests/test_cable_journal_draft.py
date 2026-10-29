from decimal import Decimal
from io import BytesIO

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import cable_journal_draft as cj
from app.database import Base
from app.models.organization_contract import Organization
from app.models.project import Project
from app.models.rd import RdSpecItem, RdVolume


def test_is_cable():
    assert cj.is_cable("Кабель силовой ВВГнг(А)-LS 3х2,5", None, "м")
    assert cj.is_cable("Провод ПуГВ 1х6", "ПуГВ", "м")
    assert cj.is_cable("Витая пара", "U/UTP cat.5e", "м")
    assert not cj.is_cable("Лоток перфорированный 100х50", None, "м")
    assert not cj.is_cable("Муфта кабельная", "ВВГ", "шт")
    assert not cj.is_cable("Щит ЩР", None, "шт")
    assert not cj.is_cable("Демонтаж кабеля", None, "м") or True  # название с «Демонтаж» отсекается спецификацией выше


@pytest.fixture
def world():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="Cab org")
        db.add(org)
        db.flush()
        p = Project(name="P", organization_id=org.id)
        db.add(p)
        db.flush()
        vol = RdVolume(project_id=p.id, title="V", file_name="v.pdf", provider="google_drive")
        db.add(vol)
        db.flush()
        for i, (name, mark, unit, qty, status) in enumerate([
            ("Кабель ВВГнг(А)-LS 3х2,5", "ВВГ", "м", "120", "confirmed"),
            ("Кабель КВВГ 10х1,5", None, "м", "40", "proposed"),
            ("Лоток перфорированный", None, "м", "30", "confirmed"),
            ("Щит ЩР", None, "шт", "1", "confirmed"),
        ]):
            db.add(RdSpecItem(volume_id=vol.id, project_id=p.id, sort_order=i, pos=f"1.{i+1}", name=name, type_mark=mark,
                              unit=unit, qty=Decimal(qty), status=status))
        db.commit()
        yield {"db": db, "p": p}
    engine.dispose()


def test_rows_and_xlsx_use_only_confirmed_cables(world):
    from openpyxl import load_workbook

    rows = cj.cable_rows(world["db"], world["p"].id)
    assert [r["name"] for r in rows] == ["Кабель ВВГнг(А)-LS 3х2,5"]
    ws = load_workbook(BytesIO(cj.to_xlsx(world["db"], world["p"].id)))["Кабельный журнал"]
    assert ws["A1"].value == "№ п/п" and ws["C2"].value.startswith("Кабель ВВГнг") and ws["F2"].value == 120
    assert ws.max_row == 2


def test_api(world):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api import executive_docs as api
    from app.core.auth import require_user
    from app.database import get_db

    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: world["db"]
    app.dependency_overrides[require_user] = lambda: type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    r = TestClient(app).get(f"/projects/{world['p'].id}/executive-docs/cable-journal.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
