from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import work_log as wl
from app.database import Base
from app.models.organization_contract import Organization
from app.models.project import Project
from app.models.rd import RdSpecItem, RdVolume


@pytest.fixture
def world():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="WL org")
        db.add(org)
        db.flush()
        p = Project(name="P", organization_id=org.id)
        o = Project(name="O", organization_id=org.id)
        db.add_all([p, o])
        db.flush()
        db.commit()
        yield {"db": db, "p": p, "o": o}
    engine.dispose()


def _add(db, pid, **kw):
    data = {"work_date": "2026-10-20", "work_name": "Прокладка кабеля", "qty": "10", "unit": "м"}
    data.update(kw)
    return wl.add_entry(db, pid, 1, data)


def _synthetic_estimate(db, pid):
    """Смета с одной позицией 100 м × 50 = 5000; возвращает (estimate_id, line_id)."""
    from app import estimates as est_svc

    parsed = {"lines": [{"sort_order": 1, "position": "1", "name": "Прокладка кабеля", "unit": "м",
                         "qty": "100", "unit_price": "50.00", "total": "5000.00"}], "total": "5000.00"}
    est = est_svc.create_estimate(db, pid, "Смета журнала", parsed)
    line_id = est_svc.get_estimate(db, pid, est["id"])["lines"][0]["id"]
    return est["id"], line_id


def _line(db, pid, est_id, line_id):
    from app import estimates as est_svc

    return next(l for l in est_svc.get_estimate(db, pid, est_id)["lines"] if l["id"] == line_id)


def test_add_list_update_delete(world):
    db, pid = world["db"], world["p"].id
    e = _add(db, pid)
    db.commit()
    assert wl.list_entries(db, pid)[0]["work_name"] == "Прокладка кабеля"
    wl.update_entry(db, pid, e["id"], {"qty": "12,5", "note": "ок"})
    assert wl.list_entries(db, pid)[0]["qty"] == "12.500"
    wl.delete_entry(db, pid, e["id"])
    assert wl.list_entries(db, pid) == []


def test_validation(world):
    db, pid = world["db"], world["p"].id
    with pytest.raises(ValueError):
        _add(db, pid, work_name="  ")
    with pytest.raises(ValueError):
        _add(db, pid, work_date=(date.today() + timedelta(days=5)).isoformat())
    with pytest.raises(ValueError):
        _add(db, pid, qty="-3")
    with pytest.raises(ValueError):
        _add(db, pid, workers=-1)
    assert _add(db, pid, qty="12,5")["qty"] == "12.500"
    vol = RdVolume(project_id=world["o"].id, title="V", file_name="v.pdf", provider="google_drive")
    db.add(vol)
    db.flush()
    foreign = RdSpecItem(volume_id=vol.id, project_id=world["o"].id, name="X", status="confirmed")
    db.add(foreign)
    db.flush()
    with pytest.raises(LookupError):
        _add(db, pid, rd_spec_item_id=foreign.id)


def test_foreign_project_and_filters(world):
    db, pid = world["db"], world["p"].id
    _add(db, pid, work_date="2026-10-01", work_name="Монтаж щита", location="Этаж 1")
    _add(db, pid, work_date="2026-10-20", work_name="Прокладка кабеля")
    e = _add(db, world["o"].id)
    db.commit()
    assert len(wl.list_entries(db, pid)) == 2
    assert [x["work_name"] for x in wl.list_entries(db, pid, date_from="2026-10-10")] == ["Прокладка кабеля"]
    assert [x["work_name"] for x in wl.list_entries(db, pid, date_to="2026-10-10")] == ["Монтаж щита"]
    assert [x["work_name"] for x in wl.list_entries(db, pid, q="ЭТАЖ")] == ["Монтаж щита"]
    with pytest.raises(LookupError):
        wl.update_entry(db, pid, e["id"], {"note": "x"})
    with pytest.raises(LookupError):
        wl.delete_entry(db, pid, e["id"])


def test_summary_suggestions_xlsx(world):
    db, pid = world["db"], world["p"].id
    _add(db, pid, qty="10", object_name="ДУБНА")
    _add(db, pid, qty="5", work_date="2026-10-21")
    vol = RdVolume(project_id=pid, title="V", file_name="v.pdf", provider="google_drive")
    db.add(vol)
    db.flush()
    db.add(RdSpecItem(volume_id=vol.id, project_id=pid, name="Насос", unit="шт", status="confirmed"))
    db.commit()
    s = wl.summary(db, pid)
    assert s["entries"] == 2 and s["days"] == 2
    assert s["by_work"][0]["qty_total"] == "15.000"
    sg = wl.suggestions(db, pid)
    assert sg["works"][0]["work_name"] == "Прокладка кабеля" and sg["objects"] == ["ДУБНА"]
    assert sg["rd_items"][0]["name"] == "Насос"
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(wl.to_xlsx(db, pid)))
    assert wb["Журнал работ"]["A1"].value == "Дата"
    assert wb["Журнал работ"].max_row == 3


def test_api(world, monkeypatch):
    from fastapi import FastAPI, HTTPException
    from fastapi.testclient import TestClient
    from app.api import work_log as api
    from app.core.auth import require_user
    from app.database import get_db

    db, pid, other = world["db"], world["p"].id, world["o"].id
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_user] = lambda: type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    client = TestClient(app)

    r = client.post(f"/projects/{pid}/work-log", json={"work_date": "2026-10-20", "work_name": "Монтаж щита", "qty": "2", "unit": "шт"})
    assert r.status_code == 200
    eid = r.json()["id"]
    assert client.post(f"/projects/{pid}/work-log", json={"work_date": "2026-10-20", "work_name": " "}).status_code == 422
    r = client.get(f"/projects/{pid}/work-log")
    assert r.status_code == 200 and len(r.json()["entries"]) == 1 and r.json()["summary"]["entries"] == 1
    assert client.get(f"/projects/{pid}/work-log/suggestions").status_code == 200
    r = client.get(f"/projects/{pid}/work-log/export.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.patch(f"/projects/{pid}/work-log/{eid}", json={"note": "ок"}).json()["note"] == "ок"

    def scoped(_db, _user, p, minimum="viewer"):
        if p != pid:
            raise HTTPException(status_code=404, detail="nf")
        return "viewer"

    monkeypatch.setattr(api, "require_project_role", scoped)
    assert client.get(f"/projects/{other}/work-log").status_code == 404
    assert client.delete(f"/projects/{pid}/work-log/{eid}").json() == {"deleted": eid}
    assert client.delete(f"/projects/{pid}/work-log/{eid}").status_code == 404


def test_schedule_overview_links_plan_and_fact(world):
    from app.models.execution_finance import ScheduleBaseline, ScheduleItem
    from app.models.user import User

    db, pid = world["db"], world["p"].id
    assert wl.schedule_overview(db, pid) == {"baseline": None, "items": []}
    user = User(email="u@x.ru", name="U")
    db.add(user)
    db.flush()
    base = ScheduleBaseline(project_id=pid, created_by_user_id=user.id, name="ГПР", version=1, status="approved")
    db.add(base)
    db.flush()
    work = ScheduleItem(project_id=pid, baseline_id=base.id, title="Прокладка кабеля", sort_order=1,
                        planned_start=date(2026, 10, 1), planned_finish=date(2026, 10, 10))
    db.add(work)
    db.flush()
    wl.add_entry(db, pid, user.id, {"work_date": "2026-10-05", "work_name": "Прокладка кабеля", "qty": "40", "unit": "м", "schedule_item_id": work.id})
    wl.add_entry(db, pid, user.id, {"work_date": "2026-10-06", "work_name": "Прокладка кабеля", "qty": "60", "unit": "м", "schedule_item_id": work.id})
    db.commit()
    res = wl.schedule_overview(db, pid)
    item = res["items"][0]
    assert res["baseline"]["name"] == "ГПР"
    assert item["fact_qty"] == "100.000" and item["fact_entries"] == 2 and item["fact_unit"] == "м"
    assert item["fact_first"] == "2026-10-05" and item["fact_last"] == "2026-10-06"
    assert item["overdue"] is True


def test_hours_and_crews(world):
    db, pid = world["db"], world["p"].id
    crew = wl.add_crew(db, pid, {"name": "Бригада 1", "foreman": "Иванов И.И.", "members": "Петров — монтажник\nСидоров — электрик"})
    assert crew["members_count"] == 3
    e = _add(db, pid, crew_id=crew["id"], hours="8,5")
    assert e["hours"] == "8.50" and e["workers"] == 3 and e["crew_id"] == crew["id"]
    with pytest.raises(ValueError):
        _add(db, pid, hours="25")
    with pytest.raises(LookupError):
        _add(db, pid, crew_id=99999)
    other = wl.add_crew(db, world["o"].id, {"name": "Чужая"})
    with pytest.raises(LookupError):
        _add(db, pid, crew_id=other["id"])
    assert wl.summary(db, pid)["man_hours"] == "25.50"  # 8.5 ч × 3 рабочих
    wl.update_crew(db, pid, crew["id"], {"foreman": "Новый"})
    assert wl.list_crews(db, pid)[0]["foreman"] == "Новый"
    with pytest.raises(ValueError):
        wl.add_crew(db, pid, {"name": " "})
    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(wl.to_xlsx(db, pid)))["Журнал работ"]
    assert [c.value for c in ws[1]][7:10] == ["Бригада", "Рабочих", "Часов"]
    assert ws["H2"].value == "Бригада 1"
    wl.delete_crew(db, pid, crew["id"])
    assert wl.list_entries(db, pid)[0]["crew_id"] is None
    with pytest.raises(LookupError):
        wl.delete_crew(db, pid, other["id"])


def test_api_crews(world):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api import work_log as api
    from app.core.auth import require_user
    from app.database import get_db

    db, pid = world["db"], world["p"].id
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_user] = lambda: type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    client = TestClient(app)
    r = client.post(f"/projects/{pid}/work-log/crews", json={"name": "Бригада А", "foreman": "Петров", "members": "Иванов\nСидоров"})
    assert r.status_code == 200 and r.json()["members_count"] == 3
    cid = r.json()["id"]
    assert client.get(f"/projects/{pid}/work-log/crews").json()["crews"][0]["name"] == "Бригада А"
    assert client.post(f"/projects/{pid}/work-log/crews", json={"name": " "}).status_code == 422
    r = client.post(f"/projects/{pid}/work-log", json={"work_date": "2026-10-20", "work_name": "Монтаж", "crew_id": cid, "hours": "7,5"})
    assert r.status_code == 200 and r.json()["hours"] == "7.50" and r.json()["workers"] == 3
    assert client.patch(f"/projects/{pid}/work-log/crews/{cid}", json={"foreman": "Новый"}).json()["foreman"] == "Новый"
    assert client.delete(f"/projects/{pid}/work-log/crews/{cid}").json() == {"deleted": cid}
    assert client.delete(f"/projects/{pid}/work-log/crews/{cid}").status_code == 404


def test_labor_timesheet_per_person(world):
    from openpyxl import load_workbook

    db, pid = world["db"], world["p"].id
    e = _add(db, pid, labor=[
        {"person_name": "Петров И.", "hours": "8", "hourly_rate": "450,50"},
        {"person_name": "Сидоров А.", "hours": "6"},
        {"person_name": "  ", "hours": "5"},  # пустое имя пропускается
    ])
    assert e["workers"] == 2
    assert e["labor"] == [
        {"person_name": "Петров И.", "hours": "8.00", "hourly_rate": "450.50"},
        {"person_name": "Сидоров А.", "hours": "6.00", "hourly_rate": None},
    ]
    assert wl.summary(db, pid)["man_hours"] == "14.00"  # табель важнее «часы × рабочие»
    with pytest.raises(ValueError):
        _add(db, pid, labor=[{"person_name": "Х", "hours": "30"}])
    with pytest.raises(ValueError):
        _add(db, pid, labor=[{"person_name": "Х", "hours": "abc"}])
    with pytest.raises(ValueError):
        _add(db, pid, labor=[{"person_name": "Х", "hours": "1", "hourly_rate": "-5"}])
    upd = wl.update_entry(db, pid, e["id"], {"labor": [{"person_name": "Петров И.", "hours": "9"}]})
    assert [r["hours"] for r in upd["labor"]] == ["9.00"]
    assert wl.list_entries(db, pid)[0]["labor"][0]["hours"] == "9.00"
    ws = load_workbook(BytesIO(wl.to_xlsx(db, pid)))["Табель"]
    assert [c.value for c in ws[1]] == ["Дата", "Работник", "Часов", "Ставка, ₽/ч", "Работа"]
    assert ws["B2"].value == "Петров И." and ws["C2"].value == 9
    wl.delete_entry(db, pid, e["id"])
    from app.models.work_log import WorkLogLabor

    assert db.query(WorkLogLabor).count() == 0


def test_journal_fact_drives_estimate_and_period_preview(world):
    from app import estimates as est_svc

    db, pid = world["db"], world["p"].id
    est_id, line_id = _synthetic_estimate(db, pid)
    db.commit()

    e = _add(db, pid, work_date="2026-10-05", qty="10", estimate_line_id=line_id)
    db.commit()
    assert e["estimate_line_id"] == line_id
    row = _line(db, pid, est_id, line_id)
    assert row["fact_qty"] == "10.000" and row["done_percent"] == 10.0

    wl.update_entry(db, pid, e["id"], {"qty": "12,5"})
    db.commit()
    row = _line(db, pid, est_id, line_id)
    assert row["fact_qty"] == "12.500" and row["done_percent"] == 12.5
    prev = est_svc.preview_act(db, pid, est_id, date(2026, 10, 1), date(2026, 10, 10))
    assert len(prev["lines"]) == 1
    assert prev["lines"][0]["qty"] == "12.500" and prev["total_ex_vat"] == "625.00"

    wl.update_entry(db, pid, e["id"], {"work_date": "2026-10-20"})
    db.commit()
    assert est_svc.preview_act(db, pid, est_id, date(2026, 10, 1), date(2026, 10, 10))["lines"] == []
    moved = est_svc.preview_act(db, pid, est_id, date(2026, 10, 11), date(2026, 10, 31))
    assert len(moved["lines"]) == 1 and moved["lines"][0]["qty"] == "12.500"

    wl.delete_entry(db, pid, e["id"])
    db.commit()
    row = _line(db, pid, est_id, line_id)
    assert row["fact_qty"] == "0.000" and row["done_percent"] == 0.0
    assert est_svc.preview_act(db, pid, est_id, date(2026, 10, 1), date(2026, 10, 10))["lines"] == []
    assert est_svc.preview_act(db, pid, est_id, date(2026, 10, 11), date(2026, 10, 31))["lines"] == []


def test_act_snapshot_survives_journal_changes(world):
    from app import estimates as est_svc

    db, pid = world["db"], world["p"].id
    est_id, line_id = _synthetic_estimate(db, pid)
    e = _add(db, pid, work_date="2026-10-05", qty="10", estimate_line_id=line_id)
    db.commit()
    act = est_svc.create_act(db, pid, est_id, date(2026, 10, 1), date(2026, 10, 31))
    db.commit()
    snapshot = est_svc.act_with_totals(db, pid, act["id"])
    snap_lines = [(l["line_id"], l["qty"], l["amount"]) for l in snapshot["lines"]]
    snap_total = snapshot["period_total"]
    assert snap_lines == [(line_id, "10.000", "500.00")] and snap_total == "500.00"

    wl.update_entry(db, pid, e["id"], {"qty": "25", "work_date": "2026-10-20"})
    db.commit()
    assert _line(db, pid, est_id, line_id)["fact_qty"] == "25.000"
    after_edit = est_svc.act_with_totals(db, pid, act["id"])
    assert [(l["line_id"], l["qty"], l["amount"]) for l in after_edit["lines"]] == snap_lines
    assert after_edit["period_total"] == snap_total

    wl.delete_entry(db, pid, e["id"])
    db.commit()
    assert _line(db, pid, est_id, line_id)["fact_qty"] == "0.000"
    after_delete = est_svc.act_with_totals(db, pid, act["id"])
    assert [(l["line_id"], l["qty"], l["amount"]) for l in after_delete["lines"]] == snap_lines
    assert after_delete["period_total"] == snap_total
