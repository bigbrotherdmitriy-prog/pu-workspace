import io
import zipfile
from xml.etree import ElementTree

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import id_handover
from app.api import id_forms as api
from app.core.auth import require_user
from app.database import Base, get_db
from app.models.executive_docs import IdParticipant, IdRequirement
from app.models.organization_contract import Organization
from app.models.project import Project
from app.models.rd import RdSpecItem, RdVolume

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _text(data: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()


def _docx_root(data: bytes):
    return ElementTree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))


def _docx_paragraphs(data: bytes) -> list[str]:
    root = _docx_root(data)
    out = []
    for p in root.iter(W + "p"):
        out.append("".join(t.text or "" for t in p.iter(W + "t")))
    return out


def _docx_tables(data: bytes) -> list[list[list[str]]]:
    root = _docx_root(data)
    tables = []
    for tbl in root.iter(W + "tbl"):
        rows = []
        for tr in tbl.iter(W + "tr"):
            cells = []
            for tc in tr.iter(W + "tc"):
                cells.append("".join(t.text or "" for t in tc.iter(W + "t")))
            rows.append(cells)
        tables.append(rows)
    return tables


def _xlsx_rows(data: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _xlsx_sheet(data: bytes):
    return load_workbook(io.BytesIO(data)).active


@pytest.fixture
def world():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="ООО Монтаж")
        db.add(org)
        db.flush()
        p = Project(name="P", organization_id=org.id)
        db.add(p)
        db.flush()
        v = RdVolume(project_id=p.id, object_name="ДУБНА", section="ЭМ", title="ЭМ", file_name="em.pdf", provider="pu_storage", external_id="1")
        db.add(v)
        db.flush()
        db.add(RdSpecItem(volume_id=v.id, project_id=p.id, name="Щит ЩР-1", type_mark="ЩРн-24", unit="шт", status="confirmed"))
        db.add(RdSpecItem(volume_id=v.id, project_id=p.id, name="Лишнее", status="rejected"))
        db.add(IdRequirement(project_id=p.id, object_name="ДУБНА", title="Журнал прокладки кабелей", doc_kind="journal", source="kit"))
        db.add(IdParticipant(project_id=p.id, object_name="ДУБНА", role="customer_control", full_name="Петров П.П.", position="инженер"))
        db.commit()
        yield {"db": db, "p": p}
    engine.dispose()


def _client(world):
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: world["db"]
    app.dependency_overrides[require_user] = lambda: type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    return TestClient(app)


def test_defaults_come_from_project(world):
    h = _client(world).get(f"/projects/{world['p'].id}/executive-docs/handover/defaults", params={"object_name": "ДУБНА"}).json()
    assert h["contractor"] == "ООО Монтаж" and h["customer_rep"] == "Петров П.П., инженер" and h["object_name"] == "ДУБНА"


def test_equipment_register_from_rd(world):
    c = _client(world)
    pid = world["p"].id
    r = c.post(f"/projects/{pid}/executive-docs/equipment-register.docx", json={"object_name": "ДУБНА", "header": {"object_name": "ДУБНА"}})
    assert r.status_code == 200
    xml = _text(r.content)
    assert "ЩРн-24" in xml and "Лишнее" not in xml
    assert c.post(f"/projects/{pid}/executive-docs/equipment-register.xlsx", json={"object_name": "ДУБНА"}).status_code == 200
    assert c.post(f"/projects/{pid}/executive-docs/equipment-register.pdf", json={}).status_code == 404


def test_id_register_lists_kit_with_status(world):
    c = _client(world)
    r = c.post(f"/projects/{world['p'].id}/executive-docs/id-register.docx", json={"object_name": "ДУБНА"})
    assert r.status_code == 200
    xml = _text(r.content)
    assert "Журнал прокладки кабелей" in xml and "нет" in xml
    assert c.post(f"/projects/{world['p'].id}/executive-docs/id-register.docx", json={}).status_code == 422


def test_trial_and_commissioning_acts(world):
    c = _client(world)
    pid = world["p"].id
    r = c.post(f"/projects/{pid}/executive-docs/trial-act.docx", json={
        "header": {"act_no": "7", "start": "01.11.2026 09:00", "end": "04.11.2026 09:00", "hours": "72"},
        "items": [{"name": "Электроснабжение", "mode": "нагрузка 60 %", "result": "pass"}]})
    assert r.status_code == 200 and "72" in _text(r.content) and "Электроснабжение" in _text(r.content)
    assert c.post(f"/projects/{pid}/executive-docs/trial-act.docx", json={"items": []}).status_code == 422
    r = c.post(f"/projects/{pid}/executive-docs/commissioning-act.docx", json={"header": {"scope": "ЭМ"}})
    assert r.status_code == 200 and "ввода в эксплуатацию" in _text(r.content)


# ---------- service-level tests (no TestClient / FastAPI lifecycle) ----------


def test_service_default_header_from_synthetic_org_and_participant(world):
    h = id_handover.default_header(world["db"], world["p"].id, "ДУБНА")
    assert h["object_name"] == "ДУБНА"
    assert h["contractor"] == "ООО Монтаж"
    assert h["customer_rep"] == "Петров П.П., инженер"
    assert "date" in h


def test_service_default_header_falls_back_to_project_name(world):
    h = id_handover.default_header(world["db"], world["p"].id, None)
    assert h["object_name"] == "P"


def test_service_equipment_rows_include_confirmed_exclude_rejected(world):
    rows = id_handover.equipment_rows(world["db"], world["p"].id, "ДУБНА")
    names = [r["name"] for r in rows]
    assert "Щит ЩР-1" in names
    assert "Лишнее" not in names
    row = next(r for r in rows if r["name"] == "Щит ЩР-1")
    assert row["type_mark"] == "ЩРн-24"
    assert row["section"] == "ЭМ"


def test_service_equipment_rows_isolated_by_object(world):
    db = world["db"]
    p = world["p"]
    v2 = RdVolume(project_id=p.id, object_name="ВОЛГА", section="ЭМ", title="ЭМ", file_name="em2.pdf", provider="pu_storage", external_id="2")
    db.add(v2)
    db.flush()
    db.add(RdSpecItem(volume_id=v2.id, project_id=p.id, name="Кабель ВВГ", type_mark="ВВГ-3х2.5", unit="м", status="confirmed"))
    db.commit()
    dubna = [r["name"] for r in id_handover.equipment_rows(db, p.id, "ДУБНА")]
    volga = [r["name"] for r in id_handover.equipment_rows(db, p.id, "ВОЛГА")]
    assert "Щит ЩР-1" in dubna and "Кабель ВВГ" not in dubna
    assert "Кабель ВВГ" in volga and "Щит ЩР-1" not in volga


def test_service_equipment_rows_isolated_by_project(world):
    db = world["db"]
    p = world["p"]
    other = Project(name="P2", organization_id=p.organization_id)
    db.add(other)
    db.flush()
    v = RdVolume(project_id=other.id, object_name="ДУБНА", section="ЭМ", title="ЭМ", file_name="em3.pdf", provider="pu_storage", external_id="3")
    db.add(v)
    db.flush()
    db.add(RdSpecItem(volume_id=v.id, project_id=other.id, name="Чужой щит", status="confirmed"))
    db.commit()
    names = [r["name"] for r in id_handover.equipment_rows(db, p.id, "ДУБНА")]
    assert "Чужой щит" not in names
    assert "Щит ЩР-1" in names


def test_service_equipment_register_docx_and_xlsx_match(world):
    db = world["db"]
    p = world["p"]
    rows = id_handover.equipment_rows(db, p.id, "ДУБНА")
    header = id_handover.default_header(db, p.id, "ДУБНА")
    docx = id_handover.render_equipment_register(header, rows)
    xlsx = id_handover.render_equipment_register_xlsx(header, rows)

    docx_paras = _docx_paragraphs(docx)
    assert any("Объект: ДУБНА" in t for t in docx_paras)
    assert any("ООО Монтаж" in t for t in docx_paras)

    docx_tables = _docx_tables(docx)
    assert docx_tables, "DOCX должен содержать таблицу ведомости"
    docx_body = docx_tables[0]
    docx_names = [r[1] for r in docx_body[1:]]
    docx_marks = [r[2] for r in docx_body[1:]]
    assert "Щит ЩР-1" in docx_names
    assert "ЩРн-24" in docx_marks
    assert "Лишнее" not in docx_names

    xlsx_rows = _xlsx_rows(xlsx)
    assert "ДУБНА" in str(xlsx_rows[0][0])
    xlsx_header = xlsx_rows[1]
    assert xlsx_header[:3] == ["№", "Наименование", "Тип, марка"]
    xlsx_names = [r[1] for r in xlsx_rows[2:]]
    xlsx_marks = [r[2] for r in xlsx_rows[2:]]
    assert "Щит ЩР-1" in xlsx_names
    assert "ЩРн-24" in xlsx_marks
    assert "Лишнее" not in xlsx_names

    assert set(docx_names) == set(xlsx_names)
    assert set(docx_marks) == set(xlsx_marks)


def test_service_equipment_register_blank_optional_no_none_text(world):
    db = world["db"]
    p = world["p"]
    v = RdVolume(project_id=p.id, object_name="МИНИМУМ", section=None, title="Том", file_name="m.pdf", provider="pu_storage", external_id="9")
    db.add(v)
    db.flush()
    db.add(RdSpecItem(volume_id=v.id, project_id=p.id, name="Прибор", type_mark=None, code=None, manufacturer=None, unit=None, qty=None, status="confirmed"))
    db.commit()
    rows = id_handover.equipment_rows(db, p.id, "МИНИМУМ")
    assert rows and rows[0]["type_mark"] == "" and rows[0]["qty"] == ""
    header = {"object_name": "МИНИМУМ"}
    docx = id_handover.render_equipment_register(header, rows)
    xlsx = id_handover.render_equipment_register_xlsx(header, rows)
    assert "None" not in _text(docx)
    for row in _xlsx_rows(xlsx):
        for cell in row:
            assert cell != "None"


def test_service_equipment_register_empty_rows_raises(world):
    with pytest.raises(ValueError):
        id_handover.render_equipment_register({}, [])
    with pytest.raises(ValueError):
        id_handover.render_equipment_register_xlsx({}, [])


def test_service_register_rows_from_kit_fixture(world):
    db = world["db"]
    p = world["p"]
    rows = id_handover.register_rows(db, p.id, "ДУБНА")
    titles = [r["title"] for r in rows]
    assert "Журнал прокладки кабелей" in titles
    row = next(r for r in rows if r["title"] == "Журнал прокладки кабелей")
    assert row["kind"] == "Журнал"
    assert row["status"] in ("есть", "черновик", "нет")
    assert isinstance(row["required"], bool)
    assert isinstance(row["files"], str)


def test_service_register_docx_and_xlsx_match(world):
    db = world["db"]
    p = world["p"]
    rows = id_handover.register_rows(db, p.id, "ДУБНА")
    header = id_handover.default_header(db, p.id, "ДУБНА")
    docx = id_handover.render_register(header, rows)
    xlsx = id_handover.render_register_xlsx(header, rows)

    docx_tables = _docx_tables(docx)
    assert docx_tables
    docx_body = docx_tables[0]
    docx_titles = [r[1] for r in docx_body[1:]]
    assert "Журнал прокладки кабелей" in docx_titles

    xlsx_rows = _xlsx_rows(xlsx)
    assert "ДУБНА" in str(xlsx_rows[0][0])
    assert xlsx_rows[1][:3] == ["№", "Документ", "Вид"]
    xlsx_titles = [r[1] for r in xlsx_rows[2:]]
    assert "Журнал прокладки кабелей" in xlsx_titles
    assert set(docx_titles) == set(xlsx_titles)

    # status / required / files согласованы между форматами
    docx_by_title = {r[1]: r for r in docx_body[1:]}
    xlsx_by_title = {r[1]: r for r in xlsx_rows[2:]}
    for title in docx_titles:
        d = docx_by_title[title]
        x = xlsx_by_title[title]
        assert d[2] == x[2]  # вид
        assert d[4].startswith(x[4])  # наличие (в DOCX может быть суффикс «(по желанию)»)
        assert ("(по желанию)" in d[4]) == (x[5] == "нет")
        assert d[5] == (x[6] or "—")


def test_service_register_blank_optional_no_none_text(world):
    db = world["db"]
    p = world["p"]
    rows = id_handover.register_rows(db, p.id, "ДУБНА")
    header = {"object_name": "ДУБНА"}
    docx = id_handover.render_register(header, rows)
    xlsx = id_handover.render_register_xlsx(header, rows)
    assert "None" not in _text(docx)
    for row in _xlsx_rows(xlsx):
        for cell in row:
            assert cell != "None"


def test_service_register_empty_rows_raises(world):
    with pytest.raises(ValueError):
        id_handover.render_register({}, [])
    with pytest.raises(ValueError):
        id_handover.render_register_xlsx({}, [])


def test_service_docx_sectpr_present_for_landscape(world):
    db = world["db"]
    p = world["p"]
    rows = id_handover.equipment_rows(db, p.id, "ДУБНА")
    docx = id_handover.render_equipment_register(id_handover.default_header(db, p.id, "ДУБНА"), rows)
    root = _docx_root(docx)
    sect = root.find(".//" + W + "sectPr")
    assert sect is not None
    pg_sz = sect.find(W + "pgSz")
    assert pg_sz is not None
    # renderer задаёт landscape=True — проверяем только наличие/ориентацию, не конкретные размеры
    orient = pg_sz.get(W + "orient")
    assert orient == "landscape"


def test_service_xlsx_print_settings_if_set(world):
    db = world["db"]
    p = world["p"]
    rows = id_handover.equipment_rows(db, p.id, "ДУБНА")
    xlsx = id_handover.render_equipment_register_xlsx(id_handover.default_header(db, p.id, "ДУБНА"), rows)
    ws = _xlsx_sheet(xlsx)
    # текущий renderer не обязан задавать print_area/orientation — проверяем только консистентность,
    # если значения заданы, они должны быть валидными строками
    if ws.print_area:
        assert isinstance(ws.print_area, str) and ws.print_area
    if ws.page_setup.orientation:
        assert ws.page_setup.orientation in ("portrait", "landscape")


def test_service_trial_act_blank_optional_no_none_text(world):
    docx = id_handover.render_trial_act({"act_no": "7"}, [{"name": "Электроснабжение"}])
    text = _text(docx)
    assert "None" not in text
    assert "Электроснабжение" in text


def test_service_commissioning_act_blank_optional_no_none_text(world):
    docx = id_handover.render_commissioning_act({"scope": "ЭМ"})
    text = _text(docx)
    assert "None" not in text
    assert "ввода в эксплуатацию" in text
