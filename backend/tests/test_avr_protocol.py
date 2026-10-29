"""Тесты протокола проверки устройства АВР."""

import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.avr_protocol import to_xlsx, validate_rows, verdict


def test_validate_rows_accepts_comma_and_skips_blank():
    rows = [
        {"scenario": "Пропадание напряжения на вводе 1",
         "expected": "Переключение на ввод 2",
         "measured_time_s": "0,5", "allowed_time_s": "1,0",
         "action_ok": "yes"},
        {"scenario": "", "expected": ""},
        {},
        {"scenario": "Ручной возврат", "action_ok": "no"},
    ]
    out = validate_rows(rows)
    assert len(out) == 2
    assert out[0]["measured_time_s"] == Decimal("0.5")
    assert out[0]["allowed_time_s"] == Decimal("1.0")
    assert out[0]["action_ok"] == "yes"
    assert out[1]["allowed_time_s"] == Decimal("1.0")
    assert out[1]["measured_time_s"] is None
    assert out[1]["action_ok"] == "no"


@pytest.mark.parametrize(
    "row, needle",
    [
        ({"scenario": "", "expected": "x"}, "Строка 1"),
        ({"scenario": "С", "measured_time_s": -1}, "Строка 1"),
        ({"scenario": "С", "allowed_time_s": -0.5}, "Строка 1"),
        ({"scenario": "С", "action_ok": "maybe"}, "Строка 1"),
    ],
)
def test_validate_rows_errors(row, needle):
    with pytest.raises(ValueError) as exc:
        validate_rows([row])
    assert needle in str(exc.value)


def test_verdict_states():
    assert verdict({
        "action_ok": "yes", "measured_time_s": Decimal("0.5"),
        "allowed_time_s": Decimal("1.0"),
    }) == "годен"
    assert verdict({
        "action_ok": "yes", "measured_time_s": Decimal("1.0"),
        "allowed_time_s": Decimal("1.0"),
    }) == "годен"
    assert verdict({
        "action_ok": "yes", "measured_time_s": Decimal("1.5"),
        "allowed_time_s": Decimal("1.0"),
    }) == "не годен"
    assert verdict({
        "action_ok": "no", "measured_time_s": Decimal("0.2"),
        "allowed_time_s": Decimal("1.0"),
    }) == "не годен"
    assert verdict({
        "action_ok": "", "measured_time_s": None,
        "allowed_time_s": Decimal("1.0"),
    }) == "нет данных"
    assert verdict({
        "action_ok": "yes", "measured_time_s": None,
        "allowed_time_s": Decimal("1.0"),
    }) == "годен"


def test_to_xlsx_layout_and_red_fill():
    header = {
        "object_name": "Объект",
        "protocol_no": "7",
        "date": "2025-01-01",
        "tested_by": "Иванов",
        "checked_by": "Петров",
        "device": "АВР-1, шкаф ШАВР-1",
        "instrument": "Секундомер СТ-1 №123",
    }
    rows = [
        {"scenario": "Пропадание напряжения на вводе 1",
         "expected": "Переключение на ввод 2",
         "measured_time_s": 0.4, "allowed_time_s": 1.0, "action_ok": "yes"},
        {"scenario": "Пропадание напряжения на вводе 2",
         "expected": "Переключение на ввод 1",
         "measured_time_s": 2.0, "allowed_time_s": 1.0, "action_ok": "yes"},
        {"scenario": "Ручной возврат",
         "expected": "Возврат на ввод 1",
         "action_ok": "no"},
    ]
    data = to_xlsx(header, rows)
    assert isinstance(data, bytes) and data[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Протокол"]

    found_title = False
    found_conclusion = False
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                if "ПРОТОКОЛ № 7" in cell.value:
                    found_title = True
                if cell.value.startswith("Заключение:"):
                    found_conclusion = True
        if any(
            isinstance(c.value, str) and c.value == "Сценарий" for c in row
        ):
            header_row = row[0].row
    assert found_title
    assert found_conclusion
    assert header_row is not None

    headers = [c.value for c in ws[header_row]]
    assert "Сценарий" in headers
    assert "Измерено, с" in headers
    assert "Допустимо, с" in headers
    assert "Заключение" in headers

    # Строка 3 — негодная, должна быть красная заливка хотя бы в одной ячейке
    red_found = False
    for r in range(header_row + 1, header_row + 1 + len(rows)):
        for j in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=j)
            if cell.fill is not None and cell.fill.fgColor.rgb in (
                "00FFC7CE", "FFC7CE",
            ):
                red_found = True
    assert red_found


class _FakeUser:
    id = 1
    username = "tester"
    is_admin = False


def _client_for(monkeypatch):
    from app.api import avr_protocol as api_mod

    app = FastAPI()
    app.include_router(api_mod.router)

    def _fake_require_user():
        return _FakeUser()

    def _fake_require_project_role(db, user, project_id, role="viewer"):
        return None

    def _fake_get_db():
        yield None

    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "require_project_role", _fake_require_project_role)
    monkeypatch.setattr(api_mod, "get_db", lambda: (yield None))
    return TestClient(app)


def test_api_returns_xlsx(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1", "object_name": "О"},
        "rows": [
            {"scenario": "Пропадание напряжения на вводе 1",
             "action_ok": "yes", "measured_time_s": 0.4},
        ],
    }
    resp = client.post(
        "/projects/1/executive-docs/avr-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert "attachment" in resp.headers.get("content-disposition", "")


def test_api_returns_422_on_error(monkeypatch):
    client = _client_for(monkeypatch)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [{"scenario": "С", "action_ok": "maybe"}],
    }
    resp = client.post(
        "/projects/1/executive-docs/avr-protocol.xlsx", json=body,
    )
    assert resp.status_code == 422
    assert "Строка 1" in resp.json()["detail"]


def test_api_viewer_sufficient(monkeypatch):
    from app.api import avr_protocol as api_mod

    calls = {}

    def _fake_get_db():
        yield None

    def _fake_require_user():
        return _FakeUser()

    def _fake_role(db, user, project_id, role="viewer"):
        calls["role"] = role
        return None

    app = FastAPI()
    app.include_router(api_mod.router)
    app.dependency_overrides[api_mod.require_user] = _fake_require_user
    monkeypatch.setattr(api_mod, "get_db", _fake_get_db)
    monkeypatch.setattr(api_mod, "require_project_role", _fake_role)

    client = TestClient(app)
    body = {
        "header": {"protocol_no": "1"},
        "rows": [{"scenario": "С", "action_ok": "yes"}],
    }
    resp = client.post(
        "/projects/1/executive-docs/avr-protocol.xlsx", json=body,
    )
    assert resp.status_code == 200
    assert calls.get("role") == "viewer"
