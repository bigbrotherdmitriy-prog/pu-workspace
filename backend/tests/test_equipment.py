from datetime import date, timedelta

import pytest
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app import equipment as equipment_service
from app.database import Base
from app.models.equipment import EquipmentUnit
from app.models.organization_contract import Organization
from app.models.project import Project
from app.models.rd import RdSpecItem, RdVolume


@pytest.fixture
def world():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_conn, _rec):
        cur = dbapi_conn.cursor()
        cur.execute("PRAGMA foreign_keys=ON")
        cur.close()

    Base.metadata.create_all(engine)
    with Session(engine) as db:
        org = Organization(name="Equip org")
        db.add(org)
        db.flush()
        project = Project(name="Equip project", organization_id=org.id)
        other = Project(name="Other", organization_id=org.id)
        db.add_all([project, other])
        db.flush()
        vol = RdVolume(
            project_id=project.id, title="Том 1", file_name="v1.pdf",
            provider="local", external_id="v1",
        )
        vol_other = RdVolume(
            project_id=other.id, title="Том 1", file_name="v1.pdf",
            provider="local", external_id="v1",
        )
        db.add_all([vol, vol_other])
        db.flush()
        db.commit()
        yield {
            "db": db,
            "project": project,
            "other": other,
            "volume": vol,
            "volume_other": vol_other,
        }
    engine.dispose()


def _make_spec(db, project_id, volume_id, name, *, unit="шт", qty="1",
               status="confirmed", type_mark=None, manufacturer=None):
    spec = RdSpecItem(
        project_id=project_id,
        volume_id=volume_id,
        name=name,
        unit=unit,
        qty=qty,
        status=status,
        type_mark=type_mark,
        manufacturer=manufacturer,
    )
    db.add(spec)
    db.flush()
    return spec


def test_generate_creates_units_and_skips_materials(world):
    db = world["db"]
    project = world["project"]
    vol = world["volume"]
    _make_spec(db, project.id, vol.id, "Щит", qty="3", unit="шт")
    _make_spec(db, project.id, vol.id, "Кабель", qty="10", unit="м")
    _make_spec(db, project.id, vol.id, "Груз", qty="5", unit="кг")
    _make_spec(db, project.id, vol.id, "Комплект", qty="2", unit="компл")
    _make_spec(db, project.id, vol.id, "К-т", qty="1", unit="к-т")
    _make_spec(db, project.id, vol.id, "Предложено", qty="2", unit="шт",
               status="proposed")
    _make_spec(db, project.id, vol.id, "Отклонено", qty="2", unit="шт",
               status="rejected")
    db.commit()

    res = equipment_service.generate(db, project.id)
    db.commit()
    assert res["created"] == 3 + 2 + 1
    units = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id
    ).all()
    assert len(units) == 6
    names = sorted(u.name for u in units)
    assert "Щит" in names and "Комплект" in names and "К-т" in names
    assert "Кабель" not in names and "Груз" not in names
    # unit_no от 1..qty
    shields = [u for u in units if u.name == "Щит"]
    assert sorted(u.unit_no for u in shields) == [1, 2, 3]


def test_generate_idempotent_and_preserves_serial(world):
    db = world["db"]
    project = world["project"]
    vol = world["volume"]
    _make_spec(db, project.id, vol.id, "Щит", qty="2", unit="шт")
    db.commit()

    res1 = equipment_service.generate(db, project.id)
    db.commit()
    assert res1["created"] == 2

    units = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id
    ).all()
    first = units[0]
    equipment_service.update_unit(
        db, project.id, first.id, {"serial_number": "SN-001"}
    )
    db.commit()

    res2 = equipment_service.generate(db, project.id)
    db.commit()
    assert res2["created"] == 0
    assert res2["existing"] == 2
    db.refresh(first)
    assert first.serial_number == "SN-001"

    # уменьшение qty не удаляет лишние юниты
    spec = db.query(RdSpecItem).filter(
        RdSpecItem.project_id == project.id
    ).first()
    spec.qty = 1
    db.commit()
    equipment_service.generate(db, project.id)
    db.commit()
    assert db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id
    ).count() == 2


def test_update_unit_serial_and_errors(world):
    db = world["db"]
    project = world["project"]
    other = world["other"]
    vol = world["volume"]
    vol_other = world["volume_other"]
    _make_spec(db, project.id, vol.id, "Щит", qty="2", unit="шт",
               type_mark="ЩР-1")
    _make_spec(db, project.id, vol.id, "Щит2", qty="1", unit="шт",
               type_mark="ЩР-2")
    _make_spec(db, other.id, vol_other.id, "Чужой", qty="1", unit="шт")
    db.commit()
    equipment_service.generate(db, project.id)
    equipment_service.generate(db, other.id)
    db.commit()

    units = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id,
        EquipmentUnit.type_mark == "ЩР-1",
    ).order_by(EquipmentUnit.unit_no).all()
    a, b = units[0], units[1]
    equipment_service.update_unit(db, project.id, a.id, {"serial_number": "SN-1"})
    db.commit()

    with pytest.raises(ValueError):
        equipment_service.update_unit(db, project.id, b.id, {"serial_number": "SN-1"})
    db.rollback()

    # разные type_mark — не конфликт
    other_unit = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id,
        EquipmentUnit.type_mark == "ЩР-2",
    ).first()
    equipment_service.update_unit(db, project.id, other_unit.id, {"serial_number": "SN-1"})
    db.commit()
    assert True

    # чужой проект -> LookupError
    foreign = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == other.id
    ).first()
    with pytest.raises(LookupError):
        equipment_service.update_unit(db, project.id, foreign.id, {"note": "x"})
    db.rollback()


def test_next_service_and_overdue(world):
    db = world["db"]
    project = world["project"]
    vol = world["volume"]
    _make_spec(db, project.id, vol.id, "Щит", qty="1", unit="шт")
    db.commit()
    equipment_service.generate(db, project.id)
    db.commit()

    unit = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id
    ).first()
    equipment_service.update_unit(
        db, project.id, unit.id,
        {"installed_date": "2024-01-15", "service_interval_months": 6},
    )
    db.commit()
    rows = equipment_service.list_units(db, project.id)
    row = rows[0]
    assert row["next_service_date"] == "2024-07-15"
    assert row["service_overdue"] is True

    # с последним ТО
    equipment_service.update_unit(db, project.id, unit.id, {
        "last_service_date": "2025-01-15", "service_interval_months": 6,
    })
    db.commit()
    row = equipment_service.list_units(db, project.id)[0]
    assert row["next_service_date"] == "2025-07-15"


def test_summary(world):
    db = world["db"]
    project = world["project"]
    vol = world["volume"]
    _make_spec(db, project.id, vol.id, "Щит", qty="3", unit="шт")
    db.commit()
    equipment_service.generate(db, project.id)
    db.commit()

    units = db.query(EquipmentUnit).filter(
        EquipmentUnit.project_id == project.id
    ).order_by(EquipmentUnit.id).all()
    equipment_service.update_unit(db, project.id, units[0].id,
                                  {"status": "received", "serial_number": "A1"})
    equipment_service.update_unit(db, project.id, units[1].id,
                                  {"status": "installed",
                                   "installed_date": "2020-01-01",
                                   "service_interval_months": 1})
    db.commit()

    s = equipment_service.summary(db, project.id)
    assert s["total"] == 3
    assert s["by_status"]["planned"] == 1
    assert s["by_status"]["received"] == 1
    assert s["by_status"]["installed"] == 1
    assert s["with_serial"] == 1
    assert s["service_overdue"] == 1


def test_to_xlsx(world):
    db = world["db"]
    project = world["project"]
    vol = world["volume"]
    _make_spec(db, project.id, vol.id, "Щит", qty="2", unit="шт")
    db.commit()
    equipment_service.generate(db, project.id)
    db.commit()

    content = equipment_service.to_xlsx(db, project.id)
    assert isinstance(content, bytes)
    assert len(content) > 0

    from openpyxl import load_workbook
    from io import BytesIO
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    assert ws.title == "Оборудование"
    headers = [c.value for c in ws[1]]
    assert headers[0] == "№"
    assert headers[1] == "Наименование"


def test_api_equipment(world, monkeypatch):
    from fastapi import FastAPI, HTTPException
    from fastapi.testclient import TestClient
    from app.api import equipment as equipment_api
    from app.core.auth import require_user
    from app.database import get_db

    db = world["db"]
    project = world["project"]
    other = world["other"]
    vol = world["volume"]
    vol_other = world["volume_other"]
    _make_spec(db, project.id, vol.id, "Щит", qty="2", unit="шт")
    _make_spec(db, other.id, vol_other.id, "Чужой", qty="1", unit="шт")
    db.commit()

    fake_user = type("U", (), {"id": 1, "is_active": True, "is_admin": True})()
    app = FastAPI()
    app.include_router(equipment_api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_user] = lambda: fake_user
    client = TestClient(app)

    # POST generate
    r = client.post(f"/projects/{project.id}/equipment/generate")
    assert r.status_code == 200
    body = r.json()
    assert body["created"] == 2

    # GET list
    r = client.get(f"/projects/{project.id}/equipment")
    assert r.status_code == 200
    data = r.json()
    assert len(data["units"]) == 2
    assert data["summary"]["total"] == 2

    unit_id = data["units"][0]["id"]

    # PATCH
    r = client.patch(
        f"/projects/{project.id}/equipment/{unit_id}",
        json={"serial_number": "SN-X", "status": "received"},
    )
    assert r.status_code == 200
    assert r.json()["serial_number"] == "SN-X"

    # 409 duplicate
    other_unit_id = data["units"][1]["id"]
    r = client.patch(
        f"/projects/{project.id}/equipment/{other_unit_id}",
        json={"serial_number": "SN-X"},
    )
    assert r.status_code == 409

    # 404 чужой проект
    def viewer_scoped(_db, _user, pid, minimum="viewer"):
        if pid != project.id:
            raise HTTPException(status_code=404, detail="not found")
        return "viewer"

    monkeypatch.setattr(equipment_api, "require_project_role", viewer_scoped)

    assert client.get(
        f"/projects/{project.id}/equipment"
    ).status_code == 200
    assert client.get(
        f"/projects/{other.id}/equipment"
    ).status_code == 404

    # export.xlsx не перехватывается /{unit_id}
    r = client.get(f"/projects/{project.id}/equipment/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    from openpyxl import load_workbook
    from io import BytesIO
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert ws.title == "Оборудование"


def test_units_of_rejected_rd_items_are_hidden(world):
    from app import equipment as eq
    from app.models.rd import RdSpecItem, RdVolume

    db, project = world["db"], world["project"]
    vol = RdVolume(project_id=project.id, title="V", file_name="v.pdf", provider="google_drive")
    db.add(vol)
    db.flush()
    keep = RdSpecItem(volume_id=vol.id, project_id=project.id, name="Насос", unit="шт", qty=1, status="confirmed")
    drop = RdSpecItem(volume_id=vol.id, project_id=project.id, name="Демонтаж насоса", unit="шт", qty=1, status="confirmed")
    db.add_all([keep, drop])
    db.flush()
    eq.generate(db, project.id)
    db.commit()
    assert len(eq.list_units(db, project.id)) == 2
    drop.status = "rejected"
    db.commit()
    names = [u["name"] for u in eq.list_units(db, project.id)]
    assert names == ["Насос"]
    assert eq.summary(db, project.id)["total"] == 1
