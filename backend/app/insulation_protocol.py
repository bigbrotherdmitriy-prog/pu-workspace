"""Протокол измерения сопротивления изоляции кабелей и проводов.

Форма по И 1.13-07 / ПУЭ гл. 1.8 (упрощённо): валидация строк измерений,
норма минимального сопротивления и генерация Excel.
"""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PAIR_ORDER: tuple[str, ...] = (
    "A-B",
    "B-C",
    "C-A",
    "A-N",
    "B-N",
    "C-N",
    "A-PE",
    "B-PE",
    "C-PE",
    "N-PE",
    "жила-оболочка",
)

ALLOWED_VOLTAGES: tuple[int, ...] = (500, 1000, 2500)

MAX_ROWS = 500

_DEFAULT_VOLTAGE = 1000


def _parse_decimal(value: Any, row_no: int, field_name: str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise ValueError(f"Строка {row_no}: поле «{field_name}» должно быть числом")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, str):
        s = value.strip().replace(",", ".")
        if not s:
            raise ValueError(f"Строка {row_no}: поле «{field_name}» пустое")
        try:
            return Decimal(s)
        except InvalidOperation:
            raise ValueError(
                f"Строка {row_no}: поле «{field_name}» должно быть числом"
            )
    raise ValueError(f"Строка {row_no}: поле «{field_name}» должно быть числом")


def _is_blank_row(raw: dict) -> bool:
    if not isinstance(raw, dict):
        return False
    line_name = raw.get("line_name")
    has_line = isinstance(line_name, str) and bool(line_name.strip())
    if has_line:
        return False
    pairs = raw.get("pairs")
    if isinstance(pairs, dict):
        for v in pairs.values():
            if v is None:
                continue
            if isinstance(v, str) and not v.strip():
                continue
            return False
    return True


def validate_rows(rows: list[dict]) -> list[dict]:
    """Проверить и нормализовать строки измерений."""
    if not isinstance(rows, list):
        raise ValueError("rows должен быть списком")
    result: list[dict] = []
    for idx, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"Строка {idx}: ожидается объект")
        if _is_blank_row(raw):
            continue
        line_name = raw.get("line_name")
        if not isinstance(line_name, str) or not line_name.strip():
            raise ValueError(f"Строка {idx}: не указано обозначение линии (line_name)")
        cable_mark_raw = raw.get("cable_mark")
        cable_mark = cable_mark_raw.strip() if isinstance(cable_mark_raw, str) else ""

        length_raw = raw.get("length_m")
        if length_raw is None or (isinstance(length_raw, str) and not length_raw.strip()):
            length_m = Decimal("0")
        else:
            length_m = _parse_decimal(length_raw, idx, "Длина")
            if length_m < 0:
                raise ValueError(f"Строка {idx}: длина не может быть отрицательной")

        voltage_raw = raw.get("voltage_v")
        if voltage_raw is None or (isinstance(voltage_raw, str) and not voltage_raw.strip()):
            voltage_v = _DEFAULT_VOLTAGE
        else:
            voltage_dec = _parse_decimal(voltage_raw, idx, "Напряжение мегаомметра")
            try:
                voltage_int = int(voltage_dec)
            except (ValueError, OverflowError):
                raise ValueError(
                    f"Строка {idx}: напряжение мегаомметра должно быть 500, 1000 или 2500 В"
                )
            if Decimal(voltage_int) != voltage_dec or voltage_int not in ALLOWED_VOLTAGES:
                raise ValueError(
                    f"Строка {idx}: напряжение мегаомметра должно быть 500, 1000 или 2500 В"
                )
            voltage_v = voltage_int

        pairs_raw = raw.get("pairs") or {}
        if not isinstance(pairs_raw, dict):
            raise ValueError(f"Строка {idx}: поле «pairs» должно быть объектом")
        pairs: dict[str, Decimal] = {}
        for key, value in pairs_raw.items():
            if key not in PAIR_ORDER:
                raise ValueError(
                    f"Строка {idx}: неизвестная пара жил «{key}»"
                )
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            pairs_dec = _parse_decimal(value, idx, f"Сопротивление {key}")
            if pairs_dec < 0:
                raise ValueError(
                    f"Строка {idx}: сопротивление пары «{key}» не может быть отрицательным"
                )
            pairs[key] = pairs_dec

        note_raw = raw.get("note")
        note = note_raw.strip() if isinstance(note_raw, str) else ""

        result.append(
            {
                "line_name": line_name.strip(),
                "cable_mark": cable_mark,
                "length_m": length_m,
                "voltage_v": voltage_v,
                "pairs": pairs,
                "note": note,
            }
        )
    if len(result) > MAX_ROWS:
        raise ValueError(f"Слишком много строк: {len(result)} (максимум {MAX_ROWS})")
    return result


def norm_min_mohm(voltage_v: int) -> Decimal:
    if voltage_v == 2500:
        return Decimal("1.0")
    return Decimal("0.5")


def verdict(row: dict) -> str:
    pairs = row.get("pairs") or {}
    if not pairs:
        return "нет данных"
    voltage_v = int(row.get("voltage_v") or _DEFAULT_VOLTAGE)
    norm = norm_min_mohm(voltage_v)
    for value in pairs.values():
        if isinstance(value, str):
            value = Decimal(value.replace(",", "."))
        if Decimal(value) < norm:
            return "не годен"
    return "годен"


def _stringify_dt(value: Any) -> str:
    if value is None:
        return ""
    s = str(value)
    if "T" in s and len(s) > 10:
        return s[:10]
    return s


_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_RED_FILL = PatternFill("solid", fgColor="FFC7CE")
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _format_decimal(value: Decimal) -> float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def to_xlsx(header: dict, rows: list[dict]) -> bytes:
    normalized = validate_rows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"

    protocol_no = _stringify_dt(header.get("protocol_no")) or "б/н"
    title = (
        f"ПРОТОКОЛ № {protocol_no} измерения сопротивления изоляции кабелей и проводов"
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = _CENTER
    ws.row_dimensions[1].height = 28

    requisites: list[tuple[str, Any]] = [
        ("Объект", header.get("object_name")),
        ("Адрес", header.get("address")),
        ("Заказчик", header.get("customer")),
        ("Подрядчик", header.get("contractor")),
        ("Лаборатория", header.get("laboratory")),
        ("Прибор", header.get("instrument")),
        ("Дата поверки прибора", header.get("instrument_check_date")),
        ("Температура, °C", header.get("temperature")),
        ("Влажность, %", header.get("humidity")),
        ("Дата", _stringify_dt(header.get("date"))),
        ("Измерения выполнил", header.get("tested_by")),
        ("Проверил", header.get("checked_by")),
    ]

    r = 3
    for label, value in requisites:
        label_cell = ws.cell(row=r, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = _LEFT_WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        value_cell = ws.cell(row=r, column=2, value="" if value is None else str(value))
        value_cell.alignment = _LEFT_WRAP
        r += 1

    r += 1
    table_start = r

    active_pairs = [
        p for p in PAIR_ORDER if any(p in row["pairs"] for row in normalized)
    ]

    fixed_headers = [
        "№ п/п",
        "Линия",
        "Кабель (марка, сечение)",
        "Длина, м",
        "Напряжение мегаомметра, В",
    ]
    tail_headers = ["Норма, МОм", "Заключение"]

    header_row = table_start
    col = 1
    for title_text in fixed_headers:
        c = ws.cell(row=header_row, column=col, value=title_text)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
        col += 1
    for pair in active_pairs:
        c = ws.cell(row=header_row, column=col, value=pair)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
        col += 1
    for title_text in tail_headers:
        c = ws.cell(row=header_row, column=col, value=title_text)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
        col += 1
    total_cols = col - 1

    r = header_row + 1
    for i, row in enumerate(normalized, start=1):
        norm = norm_min_mohm(int(row["voltage_v"]))
        row_verdict = verdict(row)

        values: list[Any] = [
            i,
            row["line_name"],
            row["cable_mark"],
            _format_decimal(row["length_m"]),
            int(row["voltage_v"]),
        ]
        pair_cols: list[int] = []
        for pair in active_pairs:
            col_idx = len(values) + 1
            pair_cols.append(col_idx)
            pair_value = row["pairs"].get(pair)
            if pair_value is None:
                values.append(None)
            else:
                values.append(_format_decimal(pair_value))
        values.append(_format_decimal(norm))
        values.append(row_verdict)

        for j, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = _BORDER
            c.alignment = _LEFT_WRAP if j in (2, 3) else _CENTER

        for pair, col_idx in zip(active_pairs, pair_cols):
            pair_value = row["pairs"].get(pair)
            if pair_value is not None and pair_value < norm:
                ws.cell(row=r, column=col_idx).fill = _RED_FILL
        r += 1

    r += 1
    bad = [row["line_name"] for row in normalized if verdict(row) == "не годен"]
    if not normalized:
        conclusion = "измерений нет"
    elif bad:
        conclusion = "есть негодные: " + ", ".join(bad)
    else:
        conclusion = "все линии годны"
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_cols)
    concl_cell = ws.cell(row=r, column=1, value=f"Заключение: {conclusion}")
    concl_cell.font = Font(bold=True)
    concl_cell.alignment = _LEFT_WRAP
    r += 2

    ws.cell(row=r, column=1, value="Измерения выполнил")
    ws.cell(row=r, column=6, value="Проверил")
    r += 1
    ws.cell(row=r, column=1, value=header.get("tested_by") or "")
    ws.cell(row=r, column=6, value=header.get("checked_by") or "")

    widths = {
        1: 6,
        2: 30,
        3: 22,
        4: 10,
        5: 16,
        total_cols - 1: 10,
        total_cols: 16,
    }
    for idx in range(6, total_cols - 1):
        widths[idx] = 12
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row_cells in ws.iter_rows(
        min_row=header_row, max_row=header_row + max(len(normalized), 0),
        min_col=1, max_col=total_cols,
    ):
        for cell in row_cells:
            if cell.alignment is None or cell.alignment.wrap_text is None:
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal if cell.alignment else None,
                    vertical=cell.alignment.vertical if cell.alignment else None,
                    wrap_text=True,
                )

    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
