"""Протокол проверки согласования параметров цепи «фаза–нуль» с защитным аппаратом.

Форма по ПУЭ 1.7.79, ГОСТ Р 50571.5.54, И 1.13-07 (упрощённо): валидация строк
измерений, расчёт тока КЗ, требуемая кратность и генерация Excel.
"""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEVICE_TYPES: tuple[str, ...] = (
    "автоматический выключатель",
    "предохранитель",
    "УЗО",
)

AUTOMAT_CHARACTERISTICS: tuple[str, ...] = ("B", "C", "D")

CHARACTERISTIC_MULTIPLIER: dict[str, Decimal] = {
    "B": Decimal("5"),
    "C": Decimal("10"),
    "D": Decimal("20"),
}

FUSE_MULTIPLIER = Decimal("3")

SAFETY_FACTOR = Decimal("1.1")

DEFAULT_PHASE_VOLTAGE = Decimal("220")

MAX_ROWS = 500


def _parse_decimal(value: Any, row_no: int, field_name: str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise ValueError(f"Строка {row_no}: поле «{field_name}» должно быть числом")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, str):
        s = value.strip().replace(",", ".")
        if not s:
            raise ValueError(f"Строка {row_no}: поле «{field_name}» пустое")
        try:
            return Decimal(s)
        except InvalidOperation:
            raise ValueError(
                f"Строка {row_no}: поле «{field_name}» должно быть числом"
            )
    raise ValueError(f"Строка {row_no}: поле «{field_name}» должно быть числом")


def _is_blank_row(raw: dict) -> bool:
    if not isinstance(raw, dict):
        return False
    for key in (
        "circuit_name",
        "device_type",
        "device_mark",
        "rated_a",
        "characteristic",
        "phase_voltage_v",
        "loop_ohm",
        "short_circuit_a",
        "note",
    ):
        v = raw.get(key)
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        return False
    return True


def _opt_decimal(raw: Any, row_no: int, field_name: str) -> Decimal | None:
    if raw is None:
        return None
    if isinstance(raw, str) and not raw.strip():
        return None
    return _parse_decimal(raw, row_no, field_name)


def validate_rows(rows: list[dict]) -> list[dict]:
    """Проверить и нормализовать строки измерений."""
    if not isinstance(rows, list):
        raise ValueError("rows должен быть списком")
    result: list[dict] = []
    for idx, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"Строка {idx}: ожидается объект")
        if _is_blank_row(raw):
            continue

        circuit_name = raw.get("circuit_name")
        if not isinstance(circuit_name, str) or not circuit_name.strip():
            raise ValueError(
                f"Строка {idx}: не указано обозначение цепи (circuit_name)"
            )

        device_type_raw = raw.get("device_type")
        if not isinstance(device_type_raw, str) or not device_type_raw.strip():
            raise ValueError(
                f"Строка {idx}: не указан тип аппарата (device_type)"
            )
        device_type = device_type_raw.strip()
        if device_type not in DEVICE_TYPES:
            raise ValueError(
                f"Строка {idx}: тип аппарата должен быть одним из: "
                "автоматический выключатель, предохранитель, УЗО"
            )

        mark_raw = raw.get("device_mark")
        device_mark = mark_raw.strip() if isinstance(mark_raw, str) else ""

        rated_raw = raw.get("rated_a")
        if rated_raw is None or (isinstance(rated_raw, str) and not rated_raw.strip()):
            raise ValueError(
                f"Строка {idx}: не указан номинальный ток аппарата (rated_a)"
            )
        rated_a = _parse_decimal(rated_raw, idx, "Номинальный ток")
        if rated_a <= 0:
            raise ValueError(
                f"Строка {idx}: номинальный ток должен быть больше нуля"
            )

        char_raw = raw.get("characteristic")
        characteristic = char_raw.strip() if isinstance(char_raw, str) else ""
        if device_type == "автоматический выключатель":
            if characteristic not in AUTOMAT_CHARACTERISTICS:
                raise ValueError(
                    f"Строка {idx}: характеристика автомата должна быть B, C или D"
                )
        elif device_type == "предохранитель":
            if characteristic:
                if characteristic not in AUTOMAT_CHARACTERISTICS:
                    raise ValueError(
                        f"Строка {idx}: характеристика предохранителя не заполняется"
                    )
            characteristic = ""
        else:
            characteristic = ""

        voltage_raw = raw.get("phase_voltage_v")
        if voltage_raw is None or (
            isinstance(voltage_raw, str) and not voltage_raw.strip()
        ):
            phase_voltage_v = DEFAULT_PHASE_VOLTAGE
        else:
            phase_voltage_v = _parse_decimal(
                voltage_raw, idx, "Фазное напряжение"
            )
            if phase_voltage_v <= 0:
                raise ValueError(
                    f"Строка {idx}: фазное напряжение должно быть больше нуля"
                )

        loop_ohm = _opt_decimal(raw.get("loop_ohm"), idx, "Сопротивление петли")
        if loop_ohm is not None and loop_ohm <= 0:
            raise ValueError(
                f"Строка {idx}: сопротивление петли должно быть больше нуля"
            )

        short_circuit_a = _opt_decimal(
            raw.get("short_circuit_a"), idx, "Ток КЗ"
        )
        if short_circuit_a is not None and short_circuit_a <= 0:
            raise ValueError(
                f"Строка {idx}: ток КЗ должен быть больше нуля"
            )

        note_raw = raw.get("note")
        note = note_raw.strip() if isinstance(note_raw, str) else ""

        result.append(
            {
                "circuit_name": circuit_name.strip(),
                "device_type": device_type,
                "device_mark": device_mark,
                "rated_a": rated_a,
                "characteristic": characteristic,
                "phase_voltage_v": phase_voltage_v,
                "loop_ohm": loop_ohm,
                "short_circuit_a": short_circuit_a,
                "note": note,
            }
        )
    if len(result) > MAX_ROWS:
        raise ValueError(f"Слишком много строк: {len(result)} (максимум {MAX_ROWS})")
    return result


def fault_current(row: dict) -> Decimal | None:
    """Ток КЗ: из loop_ohm (U/Z) или явно заданный, иначе None."""
    loop_ohm = row.get("loop_ohm")
    if loop_ohm is not None:
        if isinstance(loop_ohm, str):
            loop_ohm = Decimal(loop_ohm.replace(",", "."))
        if loop_ohm <= 0:
            return None
        voltage = row.get("phase_voltage_v") or DEFAULT_PHASE_VOLTAGE
        if isinstance(voltage, str):
            voltage = Decimal(voltage.replace(",", "."))
        return Decimal(voltage) / Decimal(loop_ohm)
    sc = row.get("short_circuit_a")
    if sc is not None:
        if isinstance(sc, str):
            sc = Decimal(sc.replace(",", "."))
        return Decimal(sc)
    return None


def required_current(row: dict) -> Decimal | None:
    """Требуемый ток КЗ с запасом 1.1 для автоматов и предохранителей."""
    device_type = row.get("device_type")
    rated = row.get("rated_a")
    if rated is None:
        return None
    if isinstance(rated, str):
        rated = Decimal(rated.replace(",", "."))
    rated = Decimal(rated)
    if device_type == "автоматический выключатель":
        characteristic = row.get("characteristic") or ""
        mult = CHARACTERISTIC_MULTIPLIER.get(characteristic)
        if mult is None:
            return None
        return SAFETY_FACTOR * mult * rated
    if device_type == "предохранитель":
        return SAFETY_FACTOR * FUSE_MULTIPLIER * rated
    return None


def verdict(row: dict) -> str:
    if fault_current(row) is None:
        return "нет данных"
    if row.get("device_type") == "УЗО":
        return "годен (проверяется кнопкой «Тест»)"
    required = required_current(row)
    if required is None:
        return "нет данных"
    ikz = fault_current(row)
    if ikz is None:
        return "нет данных"
    if ikz >= required:
        return "годен"
    return "не годен"


def _stringify_dt(value: Any) -> str:
    if value is None:
        return ""
    s = str(value)
    if "T" in s and len(s) > 10:
        return s[:10]
    return s


_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_RED_FILL = PatternFill("solid", fgColor="FFC7CE")
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _format_decimal(value: Decimal) -> float | int:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def to_xlsx(header: dict, rows: list[dict]) -> bytes:
    normalized = validate_rows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"

    protocol_no = _stringify_dt(header.get("protocol_no")) or "б/н"
    title = (
        f"ПРОТОКОЛ № {protocol_no} проверки согласования параметров цепи "
        "«фаза–нуль» с защитным аппаратом"
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = _CENTER
    ws.row_dimensions[1].height = 28

    requisites: list[tuple[str, Any]] = [
        ("Объект", header.get("object_name")),
        ("Адрес", header.get("address")),
        ("Заказчик", header.get("customer")),
        ("Подрядчик", header.get("contractor")),
        ("Лаборатория", header.get("laboratory")),
        ("Прибор", header.get("instrument")),
        ("Дата поверки прибора", header.get("instrument_check_date")),
        ("Температура, °C", header.get("temperature")),
        ("Дата", _stringify_dt(header.get("date"))),
        ("Измерения выполнил", header.get("tested_by")),
        ("Проверил", header.get("checked_by")),
    ]

    r = 3
    for label, value in requisites:
        label_cell = ws.cell(row=r, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = _LEFT_WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        value_cell = ws.cell(row=r, column=2, value="" if value is None else str(value))
        value_cell.alignment = _LEFT_WRAP
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    ref_cell = ws.cell(
        row=r,
        column=1,
        value="Справочно: допустимое время отключения при 220 В — 0.4 с "
        "(ПУЭ 1.7.79). Требуемая кратность тока КЗ к Iном — с запасом 1.1.",
    )
    ref_cell.font = Font(italic=True)
    ref_cell.alignment = _LEFT_WRAP
    r += 2

    table_start = r
    fixed_headers = [
        "№ п/п",
        "Цепь",
        "Аппарат",
        "Iном А",
        "Характеристика",
        "Uф В",
        "Zпц Ом",
        "Iкз А",
        "Требуется Iкз ≥ А",
        "Заключение",
    ]
    for col, title_text in enumerate(fixed_headers, start=1):
        c = ws.cell(row=table_start, column=col, value=title_text)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
    total_cols = len(fixed_headers)

    r = table_start + 1
    for i, row in enumerate(normalized, start=1):
        ikz = fault_current(row)
        required = required_current(row)
        row_verdict = verdict(row)

        device_label = row["device_type"]
        if row["device_mark"]:
            device_label = f"{device_label} {row['device_mark']}"

        values: list[Any] = [
            i,
            row["circuit_name"],
            device_label,
            _format_decimal(row["rated_a"]),
            row["characteristic"] or "",
            _format_decimal(row["phase_voltage_v"]),
            _format_decimal(row["loop_ohm"]) if row["loop_ohm"] is not None else None,
            _format_decimal(ikz) if ikz is not None else None,
            _format_decimal(required) if required is not None else None,
            row_verdict,
        ]

        for j, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = _BORDER
            c.alignment = _LEFT_WRAP if j in (2, 3) else _CENTER

        if row_verdict == "не годен":
            ws.cell(row=r, column=10).fill = _RED_FILL
            if ikz is not None:
                ws.cell(row=r, column=8).fill = _RED_FILL
        r += 1

    r += 1
    bad = [row["circuit_name"] for row in normalized if verdict(row) == "не годен"]
    if not normalized:
        conclusion = "измерений нет"
    elif bad:
        conclusion = "есть несоответствия: " + ", ".join(bad)
    else:
        conclusion = "все цепи соответствуют требованиям"
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_cols)
    concl_cell = ws.cell(row=r, column=1, value=f"Заключение: {conclusion}")
    concl_cell.font = Font(bold=True)
    concl_cell.alignment = _LEFT_WRAP
    r += 2

    ws.cell(row=r, column=1, value="Измерения выполнил")
    ws.cell(row=r, column=6, value="Проверил")
    r += 1
    ws.cell(row=r, column=1, value=header.get("tested_by") or "")
    ws.cell(row=r, column=6, value=header.get("checked_by") or "")

    widths = {
        1: 6,
        2: 30,
        3: 28,
        4: 10,
        5: 16,
        6: 10,
        7: 12,
        8: 12,
        9: 18,
        10: 22,
    }
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
