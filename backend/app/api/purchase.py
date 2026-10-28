from __future__ import annotations

from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.auth import require_project_role, require_user
from app.database import get_db
from app import purchase as purchase_service

router = APIRouter(prefix="/projects", tags=["purchase"])

_STATUS_LABELS = {
    "need": "Нужно заказать",
    "ordered": "Заказано",
    "paid": "Оплачено",
    "received": "Получено",
}

_HEADERS = [
    "№", "Наименование", "Тип/марка", "Производитель", "Ед.",
    "Нужно", "Заказано", "Осталось", "Статус", "Примечание",
]
_COL_WIDTHS = [6, 45, 25, 25, 8, 12, 12, 12, 18, 40]


def to_xlsx(items: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Закупка"
    ws.append(_HEADERS)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    for idx, item in enumerate(items, start=1):
        ws.append([
            idx,
            item.get("name") or "",
            item.get("type_mark") or "",
            item.get("manufacturer") or "",
            item.get("unit") or "",
            item.get("qty_needed"),
            item.get("qty_ordered"),
            item.get("qty_left"),
            _STATUS_LABELS.get(item.get("status"), item.get("status") or ""),
            item.get("note") or "",
        ])
    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class PurchaseStatusBody(BaseModel):
    status: str | None = None
    qty_ordered: str | None = None
    note: str | None = None


@router.get("/{project_id}/purchase")
def get_purchase(
    project_id: int,
    status: str | None = None,
    db: Session = Depends(get_db),
    user=Depends(require_user),
):
    require_project_role(db, user, project_id, "viewer")
    return {
        "items": purchase_service.list_items(db, project_id, status=status),
        "summary": purchase_service.summary(db, project_id),
    }


@router.post("/{project_id}/purchase/rebuild")
def rebuild_purchase(
    project_id: int,
    db: Session = Depends(get_db),
    user=Depends(require_user),
):
    require_project_role(db, user, project_id, "manager")
    return purchase_service.rebuild(db, project_id)


@router.patch("/{project_id}/purchase/{item_id}")
def patch_purchase_item(
    project_id: int,
    item_id: int,
    body: PurchaseStatusBody,
    db: Session = Depends(get_db),
    user=Depends(require_user),
):
    require_project_role(db, user, project_id, "manager")
    data = body.model_dump(exclude_unset=True)
    status = data.get("status")
    if status is None:
        raise HTTPException(status_code=422, detail="status required")
    return purchase_service.set_status(
        db,
        project_id,
        item_id,
        status,
        qty_ordered=data.get("qty_ordered"),
        note=data.get("note"),
    )


@router.get("/{project_id}/purchase/export.xlsx")
def export_purchase_xlsx(
    project_id: int,
    status: str | None = None,
    db: Session = Depends(get_db),
    user=Depends(require_user),
):
    require_project_role(db, user, project_id, "viewer")
    items = purchase_service.list_items(db, project_id, status=status)
    content = to_xlsx(items)
    return Response(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": 'attachment; filename="Zakupka.xlsx"'},
    )
