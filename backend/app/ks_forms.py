"""Печатные формы КС-2 (ОКУД 0322005) и КС-3 (ОКУД 0322001): Word и Excel."""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from app.acceptance_act_document import _MONTHS_GEN, _para, _table
from app.docx_kit import build_docx, page_break

BLANK = "________________"


def _t(v) -> str:
    return (str(v).strip() if v else "") or BLANK


def _d(v: str | None) -> str:
    if not v:
        return "«___» __________ 20__ г."
    try:
        d = date.fromisoformat(v)
    except ValueError:
        return v
    return f"«{d.day:02d}» {_MONTHS_GEN[d.month - 1]} {d.year} г."


def _short(v: str | None) -> str:
    try:
        return date.fromisoformat(v).strftime("%d.%m.%Y") if v else ""
    except ValueError:
        return v or ""


def _fmt(v: str | None) -> str:
    if v is None:
        return ""
    try:
        return f"{Decimal(v):,.2f}".replace(",", " ").replace(".", ",")
    except Exception:  # noqa: BLE001
        return str(v)


def _fmt_qty(v: str | None) -> str:
    if v is None:
        return ""
    try:
        d = Decimal(v).normalize()
        return f"{d:f}".replace(".", ",")
    except Exception:  # noqa: BLE001
        return str(v)


def ks2_rows(data: dict) -> list[list[str]]:
    rows = [["№ п/п", "Номер единичной расценки", "Наименование работ", "Номер позиции по смете", "Ед. изм.",
             "Количество", "Цена за единицу, руб.", "Стоимость, руб."]]
    for i, l in enumerate(data["lines"], start=1):
        rows.append([str(i), l["code"], l["name"], l["position"], l["unit"], _fmt_qty(l["qty"]), _fmt(l["unit_price"]), _fmt(l["amount"])])
    rows.append(["", "", "Итого", "", "", "", "", _fmt(data["period_total"])])
    return rows


def ks3_rows(data: dict) -> list[list[str]]:
    rows = [["№ п/п", "Наименование пусковых комплексов, этапов, объектов, видов работ, оборудования, затрат", "Код",
             "Стоимость выполненных работ, руб.: с начала проведения работ", "с начала года", "в том числе за отчётный период"]]
    sections = data["sections"]
    for i, (name, amount) in enumerate(sections.items(), start=1):
        rows.append([str(i), name, "", "", "", _fmt(amount)])
    rows.append(["", "Всего работ и затрат, включаемых в сметную стоимость строительства", "", _fmt(data["from_start"]), _fmt(data["from_year_start"]), _fmt(data["period_total"])])
    rows.append(["", f"в том числе НДС {_fmt_qty(data['vat_rate'])}%", "", "", "", _fmt(data["vat"])])
    rows.append(["", "Всего с НДС", "", "", "", _fmt(data["total_with_vat"])])
    return rows


def _header(header: dict, data: dict) -> list[str]:
    a = data["act"]
    return [
        f"Инвестор: {_t(header.get('investor'))}",
        f"Заказчик (генподрядчик): {_t(header.get('customer'))}",
        f"Подрядчик (субподрядчик): {_t(header.get('contractor'))}",
        f"Стройка: {_t(header.get('site'))}",
        f"Объект: {_t(header.get('object_name'))}",
        f"Договор подряда (контракт): № {_t(header.get('contract_no'))} от {_short(header.get('contract_date')) or BLANK}",
        f"Номер документа: {a['number']}    Дата составления: {_d(a['act_date'])}",
        f"Отчётный период: с {_short(a['period_from'])} по {_short(a['period_to'])}",
    ]


def render_ks2_docx(header: dict, data: dict) -> bytes:
    p: list[str] = []
    p.append(_para("Унифицированная форма № КС-2. Утверждена постановлением Госкомстата России от 11.11.99 № 100", align="right"))
    p.append(_para("Форма по ОКУД 0322005", align="right"))
    p.append(_para(""))
    for line in _header(header, data):
        p.append(_para(line))
    p.append(_para(""))
    p.append(_para("АКТ О ПРИЁМКЕ ВЫПОЛНЕННЫХ РАБОТ", bold=True, align="center"))
    p.append(_para(f"Сметная (договорная) стоимость в соответствии с договором подряда: {_fmt(data['estimate']['total'])} руб."))
    p.append(_para(""))
    p.append(_table(ks2_rows(data)))
    p.append(_para(""))
    p.append(_para(f"Всего за период без НДС: {_fmt(data['period_total'])} руб.; НДС {_fmt_qty(data['vat_rate'])}%: {_fmt(data['vat'])} руб.; всего с НДС: {_fmt(data['total_with_vat'])} руб."))
    p.append(_para(""))
    p.append(_para(f"Сдал: {_t(header.get('contractor_rep'))}  ____________ / {_t(header.get('contractor_rep_name'))} /"))
    p.append(_para(f"Принял: {_t(header.get('customer_rep'))}  ____________ / {_t(header.get('customer_rep_name'))} /"))
    return build_docx(p, landscape=True)


def render_ks3_docx(header: dict, data: dict) -> bytes:
    p: list[str] = []
    p.append(_para("Унифицированная форма № КС-3. Утверждена постановлением Госкомстата России от 11.11.99 № 100", align="right"))
    p.append(_para("Форма по ОКУД 0322001", align="right"))
    p.append(_para(""))
    for line in _header(header, data):
        p.append(_para(line))
    p.append(_para(""))
    p.append(_para("СПРАВКА О СТОИМОСТИ ВЫПОЛНЕННЫХ РАБОТ И ЗАТРАТ", bold=True, align="center"))
    p.append(_para(""))
    p.append(_table(ks3_rows(data)))
    p.append(_para(""))
    p.append(_para(f"Заказчик: {_t(header.get('customer_rep'))}  ____________ / {_t(header.get('customer_rep_name'))} /"))
    p.append(_para(f"Подрядчик: {_t(header.get('contractor_rep'))}  ____________ / {_t(header.get('contractor_rep_name'))} /"))
    return build_docx(p, landscape=True)


def render_xlsx(header: dict, data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "КС-2"
    ws.append(["Акт о приёмке выполненных работ (КС-2)"])
    ws["A1"].font = Font(bold=True, size=13)
    for line in _header(header, data):
        ws.append([line])
    ws.append([])
    start = ws.max_row + 1
    for r in ks2_rows(data):
        ws.append(r)
    for c in ws[start]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGH", (8, 22, 60, 12, 10, 12, 16, 16)):
        ws.column_dimensions[col].width = w
    ws3 = wb.create_sheet("КС-3")
    ws3.append(["Справка о стоимости выполненных работ и затрат (КС-3)"])
    ws3["A1"].font = Font(bold=True, size=13)
    for line in _header(header, data):
        ws3.append([line])
    ws3.append([])
    start3 = ws3.max_row + 1
    for r in ks3_rows(data):
        ws3.append(r)
    for c in ws3[start3]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", (8, 70, 8, 24, 16, 24)):
        ws3.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
