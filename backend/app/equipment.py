from __future__ import annotations

import re
from datetime import date, timedelta
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.equipment import EquipmentUnit
from app.models.purchase import PurchaseItemSource
from app.models.rd import RdSpecItem

_PIECE_UNITS = {"шт", "компл", "к-т", "шт."}

_STATUS_LABELS = {
    "planned": "Запланировано",
    "received": "Получено",
    "installed": "Смонтировано",
    "commissioned": "Введено в эксплуатацию",
}
_ALLOWED_STATUSES = set(_STATUS_LABELS.keys())


def _norm_unit(unit: str | None) -> str:
    if unit is None:
        return ""
    return str(unit).lower().replace(".", "").strip()


def _is_piece_unit(unit: str | None) -> bool:
    return _norm_unit(unit) in {u.replace(".", "") for u in _PIECE_UNITS}


def _qty_int(qty) -> int | None:
    if qty is None:
        return None
    try:
        value = int(qty)
    except Exception:  # noqa: BLE001
        return None
    if value < 1 or value > 200:
        return None
    # проверяем что целое (без дробной части)
    try:
        frac = qty - int(qty)
    except Exception:  # noqa: BLE001
        frac = 0
    if frac != 0:
        return None
    return value


def _add_months(d: date, months: int) -> date:
    month = d.month - 1 + months
    year = d.year + month // 12
    month = month % 12 + 1
    day = min(d.day, 28) if month == 2 else d.day
    # безопасный день: ограничим количество дней в месяце
    try:
        return d.replace(year=year, month=month, day=day)
    except ValueError:
        # последний день месяца
        if month == 12:
            next_month = date(year + 1, 1, 1)
        else:
            next_month = date(year, month + 1, 1)
        return next_month - timedelta(days=1)


def _parse_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            return None
    return None


def generate(db: Session, project_id: int) -> dict:
    specs = db.execute(
        select(RdSpecItem)
        .where(
            RdSpecItem.project_id == project_id,
            RdSpecItem.status == "confirmed",
        )
        .order_by(RdSpecItem.id)
    ).scalars().all()

    existing = db.execute(
        select(EquipmentUnit).where(EquipmentUnit.project_id == project_id)
    ).scalars().all()
    existing_keys: set[tuple[int, int]] = set()
    for unit in existing:
        if unit.rd_spec_item_id is None:
            continue
        existing_keys.add((unit.rd_spec_item_id, unit.unit_no))

    # source links: spec_item_id -> purchase_item_id
    source_rows = db.execute(
        select(PurchaseItemSource.spec_item_id, PurchaseItemSource.item_id).where(
            PurchaseItemSource.spec_item_id.in_([s.id for s in specs])
        )
    ).all() if specs else []
    purchase_by_spec: dict[int, int] = {}
    for spec_id, item_id in source_rows:
        purchase_by_spec.setdefault(spec_id, item_id)

    created = 0
    kept = 0
    for spec in specs:
        if not _is_piece_unit(spec.unit):
            continue
        qty = _qty_int(spec.qty)
        if qty is None:
            continue
        purchase_item_id = purchase_by_spec.get(spec.id)
        for unit_no in range(1, qty + 1):
            key = (spec.id, unit_no)
            if key in existing_keys:
                kept += 1
                continue
            unit = EquipmentUnit(
                project_id=project_id,
                rd_spec_item_id=spec.id,
                purchase_item_id=purchase_item_id,
                unit_no=unit_no,
                name=spec.name,
                type_mark=spec.type_mark,
                manufacturer=spec.manufacturer,
                status="planned",
            )
            db.add(unit)
            existing_keys.add(key)
            created += 1
    db.flush()
    return {"created": created, "existing": kept}


def _unit_to_dict(unit: EquipmentUnit, today: date | None = None) -> dict:
    if today is None:
        today = date.today()
    next_service: date | None = None
    if unit.service_interval_months:
        base = unit.last_service_date or unit.installed_date
        if base is not None:
            next_service = _add_months(base, unit.service_interval_months)
    return {
        "id": unit.id,
        "rd_spec_item_id": unit.rd_spec_item_id,
        "purchase_item_id": unit.purchase_item_id,
        "unit_no": unit.unit_no,
        "name": unit.name,
        "type_mark": unit.type_mark,
        "manufacturer": unit.manufacturer,
        "serial_number": unit.serial_number,
        "location": unit.location,
        "status": unit.status,
        "received_date": unit.received_date.isoformat() if unit.received_date else None,
        "installed_date": unit.installed_date.isoformat() if unit.installed_date else None,
        "warranty_until": unit.warranty_until.isoformat() if unit.warranty_until else None,
        "service_interval_months": unit.service_interval_months,
        "last_service_date": unit.last_service_date.isoformat() if unit.last_service_date else None,
        "next_service_date": next_service.isoformat() if next_service else None,
        "service_overdue": bool(next_service is not None and next_service < today),
        "note": unit.note,
    }


def _active_units_stmt(project_id: int):
    """Единицы проекта, кроме тех, чья позиция РД позже отклонена."""
    rejected = select(RdSpecItem.id).where(
        RdSpecItem.project_id == project_id, RdSpecItem.status != "confirmed",
    )
    return select(EquipmentUnit).where(
        EquipmentUnit.project_id == project_id,
        EquipmentUnit.rd_spec_item_id.not_in(rejected) | EquipmentUnit.rd_spec_item_id.is_(None),
    )


def list_units(
    db: Session,
    project_id: int,
    status: str | None = None,
    q: str | None = None,
) -> list[dict]:
    stmt = _active_units_stmt(project_id)
    if status:
        stmt = stmt.where(EquipmentUnit.status == status)
    stmt = stmt.order_by(
        EquipmentUnit.name, EquipmentUnit.type_mark, EquipmentUnit.unit_no,
    )
    units = db.execute(stmt).scalars().all()
    today = date.today()
    result: list[dict] = []
    needle = q.lower().strip() if q else None
    for unit in units:
        if needle:
            hay = " ".join(filter(None, [
                unit.name or "",
                unit.type_mark or "",
                unit.serial_number or "",
            ])).lower()
            if needle not in hay:
                continue
        result.append(_unit_to_dict(unit, today=today))
    return result


def _get_unit(db: Session, project_id: int, unit_id: int) -> EquipmentUnit:
    unit = db.execute(
        select(EquipmentUnit).where(
            EquipmentUnit.id == unit_id,
            EquipmentUnit.project_id == project_id,
        )
    ).scalar_one_or_none()
    if unit is None:
        raise LookupError("unit not found")
    return unit


def update_unit(db: Session, project_id: int, unit_id: int, data: dict) -> dict:
    unit = _get_unit(db, project_id, unit_id)

    if "serial_number" in data:
        raw = data["serial_number"]
        if isinstance(raw, str):
            value = raw.strip() or None
        else:
            value = raw
        if value:
            conflict = db.execute(
                select(EquipmentUnit).where(
                    EquipmentUnit.project_id == project_id,
                    EquipmentUnit.id != unit.id,
                    EquipmentUnit.serial_number == value,
                    EquipmentUnit.type_mark == unit.type_mark,
                )
            ).scalar_one_or_none()
            if conflict is not None:
                raise ValueError("duplicate serial")
        unit.serial_number = value

    if "location" in data:
        raw = data["location"]
        unit.location = raw.strip() if isinstance(raw, str) and raw.strip() else (None if raw is None else raw)

    if "status" in data:
        value = data["status"]
        if value not in _ALLOWED_STATUSES:
            raise ValueError("invalid status")
        unit.status = value

    if "received_date" in data:
        unit.received_date = _parse_date(data["received_date"])
    if "installed_date" in data:
        unit.installed_date = _parse_date(data["installed_date"])
    if "warranty_until" in data:
        unit.warranty_until = _parse_date(data["warranty_until"])
    if "last_service_date" in data:
        unit.last_service_date = _parse_date(data["last_service_date"])

    if "service_interval_months" in data:
        raw = data["service_interval_months"]
        if raw is None or raw == "":
            unit.service_interval_months = None
        else:
            try:
                months = int(raw)
            except Exception:  # noqa: BLE001
                raise ValueError("invalid service interval") from None
            if months < 1 or months > 120:
                raise ValueError("invalid service interval")
            unit.service_interval_months = months

    if "note" in data:
        raw = data["note"]
        if isinstance(raw, str):
            unit.note = raw.strip() or None
        else:
            unit.note = raw

    db.flush()
    return _unit_to_dict(unit)


def summary(db: Session, project_id: int) -> dict:
    units = db.execute(_active_units_stmt(project_id)).scalars().all()
    result = {
        "total": 0,
        "by_status": {k: 0 for k in _ALLOWED_STATUSES},
        "with_serial": 0,
        "service_overdue": 0,
    }
    today = date.today()
    for unit in units:
        result["total"] += 1
        if unit.status in result["by_status"]:
            result["by_status"][unit.status] += 1
        if unit.serial_number:
            result["with_serial"] += 1
        if unit.service_interval_months:
            base = unit.last_service_date or unit.installed_date
            if base is not None:
                next_service = _add_months(base, unit.service_interval_months)
                if next_service < today:
                    result["service_overdue"] += 1
    return result


_HEADERS = [
    "№", "Наименование", "Тип/марка", "Изготовитель", "Заводской №", "Место",
    "Статус", "Получено", "Смонтировано", "Гарантия до", "Интервал ТО (мес.)",
    "Последнее ТО", "Следующее ТО", "Примечание",
]
_COL_WIDTHS = [6, 40, 25, 25, 20, 20, 22, 12, 14, 12, 16, 14, 14, 40]


def _fmt(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


def to_xlsx(db: Session, project_id: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Оборудование"
    ws.append(_HEADERS)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold

    units = db.execute(
        _active_units_stmt(project_id)
        .order_by(
            EquipmentUnit.name, EquipmentUnit.type_mark, EquipmentUnit.unit_no,
        )
    ).scalars().all()

    today = date.today()
    for idx, unit in enumerate(units, start=1):
        next_service: date | None = None
        if unit.service_interval_months:
            base = unit.last_service_date or unit.installed_date
            if base is not None:
                next_service = _add_months(base, unit.service_interval_months)
        _ = today
        ws.append([
            idx,
            unit.name or "",
            unit.type_mark or "",
            unit.manufacturer or "",
            unit.serial_number or "",
            unit.location or "",
            _STATUS_LABELS.get(unit.status, unit.status or ""),
            _fmt(unit.received_date),
            _fmt(unit.installed_date),
            _fmt(unit.warranty_until),
            unit.service_interval_months if unit.service_interval_months else "",
            _fmt(unit.last_service_date),
            _fmt(next_service),
            unit.note or "",
        ])

    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
