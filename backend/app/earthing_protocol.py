"""Протокол измерения сопротивления заземляющего устройства.

Форма по ПУЭ гл. 1.7 / И 1.13-07 (упрощённо): валидация строк измерений,
норма сопротивления по напряжению сети и генерация Excel.
"""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ALLOWED_VOLTAGES: tuple[int, ...] = (380, 220, 660)

MAX_ROWS = 500

_DEFAULT_VOLTAGE = 380

_NORMS_BY_VOLTAGE: dict[int, Decimal] = {
    380: Decimal("4"),
    220: Decimal("8"),
    660: Decimal("2"),
}

_TWO_PLACES = Decimal("0.01")


def _parse_decimal(value: Any, row_no: int, field_name: str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise ValueError(f"Строка {row_no}: поле «{field_name}» должно быть числом")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, str):
        s = value.strip().replace(",", ".")
        if not s:
            raise ValueError(f"Строка {row_no}: поле «{field_name}» пустое")
        try:
            return Decimal(s)
        except InvalidOperation:
            raise ValueError(
                f"Строка {row_no}: поле «{field_name}» должно быть числом"
            )
    raise ValueError(f"Строка {row_no}: поле «{field_name}» должно быть числом")


def _is_blank_row(raw: dict) -> bool:
    if not isinstance(raw, dict):
        return False
    point_name = raw.get("point_name")
    if isinstance(point_name, str) and point_name.strip():
        return False
    measured = raw.get("measured_ohm")
    if measured is None:
        return True
    if isinstance(measured, str) and not measured.strip():
        return True
    return False


def validate_rows(rows: list[dict]) -> list[dict]:
    """Проверить и нормализовать строки измерений."""
    if not isinstance(rows, list):
        raise ValueError("rows должен быть списком")
    result: list[dict] = []
    for idx, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"Строка {idx}: ожидается объект")
        if _is_blank_row(raw):
            continue
        point_name_raw = raw.get("point_name")
        if not isinstance(point_name_raw, str) or not point_name_raw.strip():
            raise ValueError(f"Строка {idx}: не указано заземляющее устройство (point_name)")
        point_name = point_name_raw.strip()

        location_raw = raw.get("location")
        location = location_raw.strip() if isinstance(location_raw, str) else ""

        soil_raw = raw.get("soil")
        soil = soil_raw.strip() if isinstance(soil_raw, str) else ""

        electrode_raw = raw.get("electrode")
        electrode = electrode_raw.strip() if isinstance(electrode_raw, str) else ""

        network_raw = raw.get("network_voltage_v")
        if network_raw is None or (isinstance(network_raw, str) and not network_raw.strip()):
            network_voltage_v = _DEFAULT_VOLTAGE
        else:
            network_dec = _parse_decimal(network_raw, idx, "Напряжение сети")
            try:
                network_int = int(network_dec)
            except (ValueError, OverflowError):
                raise ValueError(
                    f"Строка {idx}: напряжение сети должно быть 220, 380 или 660 В"
                )
            if Decimal(network_int) != network_dec or network_int not in ALLOWED_VOLTAGES:
                raise ValueError(
                    f"Строка {idx}: напряжение сети должно быть 220, 380 или 660 В"
                )
            network_voltage_v = network_int

        measured_raw = raw.get("measured_ohm")
        measured_ohm: Decimal | None
        if measured_raw is None or (
            isinstance(measured_raw, str) and not measured_raw.strip()
        ):
            measured_ohm = None
        else:
            measured_ohm = _parse_decimal(measured_raw, idx, "Измеренное сопротивление")
            if measured_ohm < 0:
                raise ValueError(
                    f"Строка {idx}: измеренное сопротивление не может быть отрицательным"
                )
            measured_ohm = measured_ohm.quantize(_TWO_PLACES)

        season_raw = raw.get("season_factor")
        if season_raw is None or (
            isinstance(season_raw, str) and not season_raw.strip()
        ):
            season_factor = Decimal("1.0")
        else:
            season_factor = _parse_decimal(season_raw, idx, "Коэффициент сезонности")
            if season_factor < 1:
                raise ValueError(
                    f"Строка {idx}: коэффициент сезонности не может быть меньше 1"
                )

        norm_raw = raw.get("norm_ohm")
        norm_override: Decimal | None
        if norm_raw is None or (isinstance(norm_raw, str) and not norm_raw.strip()):
            norm_override = None
        else:
            norm_override = _parse_decimal(norm_raw, idx, "Норма")
            if norm_override < 0:
                raise ValueError(
                    f"Строка {idx}: норма сопротивления не может быть отрицательной"
                )

        note_raw = raw.get("note")
        note = note_raw.strip() if isinstance(note_raw, str) else ""

        result.append(
            {
                "point_name": point_name,
                "location": location,
                "soil": soil,
                "electrode": electrode,
                "network_voltage_v": network_voltage_v,
                "measured_ohm": measured_ohm,
                "season_factor": season_factor,
                "norm_override": norm_override,
                "note": note,
            }
        )
    if len(result) > MAX_ROWS:
        raise ValueError(f"Слишком много строк: {len(result)} (максимум {MAX_ROWS})")
    return result


def norm_ohm(network_voltage_v: int) -> Decimal:
    try:
        return _NORMS_BY_VOLTAGE[int(network_voltage_v)]
    except KeyError:
        raise ValueError(
            f"недопустимое напряжение сети: {network_voltage_v} "
            "(допускается 220, 380 или 660 В)"
        )


def reduced_ohm(row: dict) -> Decimal | None:
    measured = row.get("measured_ohm")
    if measured is None:
        return None
    if isinstance(measured, str):
        measured = Decimal(measured.replace(",", "."))
    season = row.get("season_factor") or Decimal("1.0")
    if isinstance(season, str):
        season = Decimal(season.replace(",", "."))
    return (Decimal(measured) * Decimal(season)).quantize(_TWO_PLACES)


def row_norm(row: dict) -> Decimal:
    override = row.get("norm_override")
    if override is not None:
        if isinstance(override, str):
            override = Decimal(override.replace(",", "."))
        return Decimal(override)
    return norm_ohm(int(row.get("network_voltage_v") or _DEFAULT_VOLTAGE))


def verdict(row: dict) -> str:
    if row.get("measured_ohm") is None:
        return "нет данных"
    value = reduced_ohm(row)
    if value is None:
        return "нет данных"
    norm = row_norm(row)
    if value <= norm:
        return "годен"
    return "не годен"


def _stringify_dt(value: Any) -> str:
    if value is None:
        return ""
    s = str(value)
    if "T" in s and len(s) > 10:
        return s[:10]
    return s


_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_RED_FILL = PatternFill("solid", fgColor="FFC7CE")
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _format_decimal(value: Decimal) -> float | int:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def to_xlsx(header: dict, rows: list[dict]) -> bytes:
    normalized = validate_rows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"

    protocol_no = _stringify_dt(header.get("protocol_no")) or "б/н"
    title = (
        f"ПРОТОКОЛ № {protocol_no} измерения сопротивления "
        "заземляющего устройства"
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = _CENTER
    ws.row_dimensions[1].height = 28

    requisites: list[tuple[str, Any]] = [
        ("Объект", header.get("object_name")),
        ("Адрес", header.get("address")),
        ("Заказчик", header.get("customer")),
        ("Подрядчик", header.get("contractor")),
        ("Лаборатория", header.get("laboratory")),
        ("Прибор", header.get("instrument")),
        ("Дата поверки прибора", header.get("instrument_check_date")),
        ("Температура, °C", header.get("temperature")),
        ("Состояние грунта", header.get("soil_state")),
        ("Дата", _stringify_dt(header.get("date"))),
        ("Измерения выполнил", header.get("tested_by")),
        ("Проверил", header.get("checked_by")),
    ]

    r = 3
    for label, value in requisites:
        label_cell = ws.cell(row=r, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = _LEFT_WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
        value_cell = ws.cell(row=r, column=2, value="" if value is None else str(value))
        value_cell.alignment = _LEFT_WRAP
        r += 1

    r += 1
    table_start = r

    headers: list[str] = [
        "№ п/п",
        "Точка/ЗУ",
        "Место",
        "Грунт",
        "Заземлитель",
        "Напряжение сети В",
        "Измерено Ом",
        "Коэф. сезонности",
        "Приведённое Ом",
        "Норма Ом",
        "Заключение",
    ]
    total_cols = len(headers)

    header_row = table_start
    for col, title_text in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=title_text)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER

    r = header_row + 1
    for i, row in enumerate(normalized, start=1):
        norm_val = row_norm(row)
        reduced = reduced_ohm(row)
        row_verdict = verdict(row)

        measured_value: Any
        if row["measured_ohm"] is None:
            measured_value = None
        else:
            measured_value = _format_decimal(row["measured_ohm"])

        reduced_value: Any
        if reduced is None:
            reduced_value = None
        else:
            reduced_value = _format_decimal(reduced)

        values: list[Any] = [
            i,
            row["point_name"],
            row["location"],
            row["soil"],
            row["electrode"],
            int(row["network_voltage_v"]),
            measured_value,
            _format_decimal(Decimal(row["season_factor"])),
            reduced_value,
            _format_decimal(norm_val),
            row_verdict,
        ]

        for j, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = _BORDER
            c.alignment = _LEFT_WRAP if j in (2, 3, 4, 5) else _CENTER

        if row_verdict == "не годен":
            for j in (7, 9):
                cell = ws.cell(row=r, column=j)
                if cell.value is not None:
                    cell.fill = _RED_FILL
        r += 1

    r += 1
    bad = [row["point_name"] for row in normalized if verdict(row) == "не годен"]
    no_data = [row["point_name"] for row in normalized if verdict(row) == "нет данных"]
    if not normalized:
        conclusion = "измерений нет"
    elif bad:
        conclusion = "есть негодные: " + ", ".join(bad)
    elif no_data:
        conclusion = "нет данных: " + ", ".join(no_data)
    else:
        conclusion = "все заземляющие устройства годны"
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_cols)
    concl_cell = ws.cell(row=r, column=1, value=f"Заключение: {conclusion}")
    concl_cell.font = Font(bold=True)
    concl_cell.alignment = _LEFT_WRAP
    r += 2

    ws.cell(row=r, column=1, value="Измерения выполнил")
    ws.cell(row=r, column=6, value="Проверил")
    r += 1
    ws.cell(row=r, column=1, value=header.get("tested_by") or "")
    ws.cell(row=r, column=6, value=header.get("checked_by") or "")

    widths = {
        1: 6,
        2: 26,
        3: 22,
        4: 16,
        5: 20,
        6: 12,
        7: 12,
        8: 14,
        9: 14,
        10: 10,
        11: 14,
    }
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row_cells in ws.iter_rows(
        min_row=header_row,
        max_row=header_row + max(len(normalized), 0),
        min_col=1,
        max_col=total_cols,
    ):
        for cell in row_cells:
            if cell.alignment is None or cell.alignment.wrap_text is None:
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal if cell.alignment else None,
                    vertical=cell.alignment.vertical if cell.alignment else None,
                    wrap_text=True,
                )

    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
